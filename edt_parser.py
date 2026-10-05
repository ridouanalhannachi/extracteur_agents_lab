
import io
import os
import re
import csv
import shutil
import subprocess
import tempfile
import unicodedata
from collections import OrderedDict
from urllib.parse import unquote

import fitz
import pandas as pd
import xlsxwriter

DAYS = {"Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"}

OUTPUT_COLUMNS = [
    "Nom et prénom",
    "Statut",
    "Filière",
    "Niveau",
    "Matière",
    "CH",
    "Tél",
    "Email",
    "Département",
]

DETAIL_COLUMNS = [
    "Jour",
    "Matière",
    "Type",
    "Nom et prénom",
    "Horaire",
    "Durée",
    "Groupe",
    "Salle",
    "Filière",
    "Niveau",
    "Année universitaire",
    "Source PDF",
    "Page",
]


def _clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _ascii_key(value):
    s = _clean(value).replace("’", "'").upper()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"^(?:PROF|PR|DR)\.?\s*", "", s)
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _normalize_colname(value):
    s = _ascii_key(value)
    return s.lower()


def parse_duration_hours(value):
    s = _clean(value).lower().replace(" ", "")
    s = s.replace("–", "-").replace("—", "-").replace("−", "-")
    m = re.search(r"(?<!\d)(\d{1,2})(?:[h:.](\d{0,2}))?-(\d{1,2})(?:[h:.](\d{0,2}))?(?!\d)", s)
    if not m:
        return None
    h1, m1 = int(m.group(1)), int(m.group(2) or 0)
    h2, m2 = int(m.group(3)), int(m.group(4) or 0)
    minutes = (h2 * 60 + m2) - (h1 * 60 + m1)
    if minutes <= 0 or minutes > 6 * 60:
        return None
    return round(minutes / 60, 2)


def extract_metadata(page, source_name=""):
    text = page.get_text("text")
    source_name = unquote(source_name)

    filiere_code = ""
    filiere_label = ""
    m = re.search(r"Fili[eè]re\s*:\s*(.+?)\s*\(([A-Z0-9_-]+)\)", text, re.I)
    if m:
        filiere_label = _clean(m.group(1))
        filiere_code = _clean(m.group(2)).upper()
    else:
        m = re.search(r"Fili[eè]re\s*:\s*(.+)", text, re.I)
        if m:
            filiere_label = _clean(m.group(1))

    if not filiere_code:
        # The new templates print the code on a separate line ("G C CD", "RT").
        for line in text.splitlines()[:28]:
            compact = re.sub(r"\s+", "", line).upper()
            if compact in {"GCCD", "RT", "IDSD", "ILCS", "FBTD", "MLT", "DAWM", "GEE", "GITAM", "IAID", "MGPDI", "WM"}:
                filiere_code = compact
                break
    if not filiere_code:
        m = re.search(r"Fili[eè]re\s*:\s*([A-Z]{2,6})\s*[–-]\s*Semestre", text, re.I)
        if m:
            filiere_code = m.group(1).upper()
    if not filiere_code:
        for code in ("GITAM", "MGPDI", "GCCD", "IDSD", "ILCS", "FBTD", "DAWM", "GEE", "IAID", "MLT", "WM", "RT"):
            if re.search(rf"(?<![A-Z]){code}(?![A-Z])", source_name.upper()):
                filiere_code = code
                break

    niveau = ""
    m = re.search(r"du\s+(\d+)\s*(?:ème|eme|e)?\s*Semestre", text, re.I)
    if m:
        niveau = f"S{m.group(1)}"
    else:
        m = (re.search(r"\bSemestre\s*([1-6])\b", text, re.I)
             or re.search(r"\bS\.?([1-6])\b", text[:1000], re.I)
             or re.search(r"\bS\.?([1-6])\b", source_name, re.I))
        if m:
            niveau = f"S{m.group(1)}"

    annee = ""
    m = re.search(
        r"Ann[ée]e\s+universitaire\s*:\s*([0-9]{4}\s*/\s*[0-9]{4})",
        text,
        re.I,
    )
    if m:
        annee = m.group(1).replace(" ", "")
    else:
        m = re.search(r"Ann[ée]e\s+universitaire\s*:?\s*([0-9]{4}\s*/\s*[0-9]{4})", text, re.I)
        if m:
            annee = m.group(1).replace(" ", "")

    return {
        "filiere": filiere_code,
        "filiere_label": filiere_label,
        "niveau": niveau,
        "annee": annee,
    }


def _best_timetable_table(page):
    tables = page.find_tables().tables
    if not tables:
        return None

    candidates = []
    for table in tables:
        data = table.extract()
        weekday_hits = sum(
            1 for row in data
            if row and _day(row[0])
        )
        x0, y0, x1, y1 = table.bbox
        area = (x1 - x0) * (y1 - y0)
        candidates.append((weekday_hits, area, table))

    candidates.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return candidates[0][2] if candidates and candidates[0][0] > 0 else None


TIME_RE = re.compile(r"(?<!\d)\d{1,2}(?:[h:.]\d{0,2})?\s*[-–—−]\s*\d{1,2}(?:[h:.]\d{0,2})?(?!\d)", re.I)
TEACHER_RE = re.compile(r"^(?:Prof\.?|Pr\.?|Dr\.?)\s+", re.I)
ROOM_RE = re.compile(r"\b(?:Salle|Amphi|SM\s*\d+|SC\s*\d+|TP\s*\d+)\b", re.I)


def _day(value):
    key = _ascii_key(value)
    for name in DAYS:
        if key == _ascii_key(name):
            return name
    return ""


def _record(day, subject, teacher, schedule, meta, source_name, page_no,
            teaching_type="", group="", room=""):
    teacher = _clean(teacher)
    if TEACHER_RE.match(teacher):
        teacher = re.sub(r"^(Prof|Pr|Dr)\.?\s+", lambda m: m.group(1).title() + ". ", teacher, flags=re.I)
    return {
        "Jour": day, "Matière": _clean(subject), "Type": _clean(teaching_type),
        "Nom et prénom": teacher, "Horaire": _clean(schedule),
        "Durée": parse_duration_hours(schedule), "Groupe": _clean(group),
        "Salle": _clean(room), "Filière": meta["filiere"],
        "Niveau": meta["niveau"], "Année universitaire": meta["annee"],
        "Source PDF": source_name, "Page": page_no,
    }


def _parse_course_fragments(raw_fragments, allow_missing_teacher=False):
    """Classify PDF/Word cell lines regardless of room/time ordering."""
    raw_lines = [_clean(piece) for raw in raw_fragments for piece in str(raw).splitlines() if _clean(piece)]
    fragments = []
    for piece in raw_lines:
        # PDF extraction sometimes wraps "Pr" onto its own line, followed by
        # the unassigned-name dots on one or two further lines.
        if fragments and fragments[-1].upper() in ("PR", "PROF", "DR"):
            fragments[-1] += " " + piece
        elif fragments and TEACHER_RE.match(fragments[-1]) and not _ascii_key(piece):
            fragments[-1] += piece
        else:
            fragments.append(piece)
    t = next((i for i, value in enumerate(fragments) if TIME_RE.search(value)), None)
    if t is None:
        return None
    schedule = TIME_RE.search(fragments[t]).group()
    if parse_duration_hours(schedule) is None:
        return None
    types = {"CM", "TD", "TP", "TDM", "COURS", "COURS/TD", "COURS/TD/TP"}
    kind = _clean(fragments[t].split(schedule, 1)[0].strip(" ·"))
    kind = kind or next((value for value in fragments if value.upper() in types), "")
    teacher_idx = next((i for i, value in enumerate(fragments) if TEACHER_RE.match(value)), None)
    room_idx = next((i for i, value in enumerate(fragments) if ROOM_RE.search(value) and not TEACHER_RE.match(value)), None)
    if teacher_idx is None:
        possibilities = [i for i in range(t) if i != room_idx and fragments[i].upper() not in types]
        if len(possibilities) >= 2:
            teacher_idx = possibilities[-1]
    if teacher_idx is None and not allow_missing_teacher:
        return None
    teacher = fragments[teacher_idx] if teacher_idx is not None else "À compléter"
    room = fragments[room_idx] if room_idx is not None else ""
    if "·" in teacher:
        person, sep, location = teacher.partition("·")
        if ROOM_RE.search(location):
            teacher, room = person.strip(), location.strip()
    subject = " ".join(value for i, value in enumerate(fragments)
                       if i < min(t, teacher_idx if teacher_idx is not None else t)
                       and i != room_idx and value.upper() not in types
                       and value not in ("—", "-"))
    group = next((value for value in fragments if re.match(r"^(?:Groupe\s+\w+|TC(?:\s|$))", value, re.I)), "")
    if not subject or _ascii_key(subject) == "EXAMEN":
        return None
    return subject, teacher, schedule, kind, group, room


def _simple_table_sessions(table, meta, source_name, page_no):
    """The FBTD/MLT grid uses subject / teacher / time / room rows."""
    rows = table.extract()
    out = []
    for i, row in enumerate(rows):
        day = _day(row[0] if row else "")
        if not day:
            continue
        stop = next((j for j in range(i + 1, len(rows)) if _day(rows[j][0] if rows[j] else "")), len(rows))
        for col in range(1, max(map(len, rows[i:stop]), default=0)):
            parsed = _parse_course_fragments(r[col] for r in rows[i:stop] if col < len(r) and _clean(r[col]))
            if parsed:
                subject, teacher, schedule, kind, group, room = parsed
                out.append(_record(day, subject, teacher, schedule, meta, source_name,
                                   page_no, kind, group, room))
    return out


def _generic_grid_sessions(rows, meta, source_name, page_no):
    """Read grids with days in rows or columns and arbitrary time headers."""
    rows = [[str(cell).strip() if cell else "" for cell in row] for row in rows]
    if not rows:
        return []
    width = max(map(len, rows))
    rows = [row + [""] * (width - len(row)) for row in rows]
    result = []
    # Transposed timetable: weekdays are column headings; the first column
    # gives the time range for each row.
    header = next((row for row in rows[:5] if sum(bool(_day(value)) for value in row) >= 2), None)
    if header is not None:
        for row in rows[rows.index(header) + 1:]:
            schedule = next((TIME_RE.search(value).group() for value in row[:2]
                             if TIME_RE.search(value)), "")
            if not schedule:
                continue
            for col, day_label in enumerate(header):
                day = _day(day_label)
                if not day or col >= len(row) or not row[col]:
                    continue
                parsed = _parse_course_fragments([row[col], schedule], allow_missing_teacher=True)
                if parsed:
                    subject, teacher, time, kind, group, room = parsed
                    result.append(_record(day, subject, teacher, time, meta,
                                          source_name, page_no, kind, group, room))
        return result
    # Conventional timetable: each day starts a row, while a course can
    # occupy one multiline cell or several subsequent table rows.
    day_rows = [(i, next(((j, _day(value)) for j, value in enumerate(row) if _day(value)), None))
                for i, row in enumerate(rows)]
    day_rows = [(i, j, day) for i, pair in day_rows if pair for j, day in [pair]]
    if day_rows:
        for pos, (start, day_col, day) in enumerate(day_rows):
            end = day_rows[pos + 1][0] if pos + 1 < len(day_rows) else len(rows)
            for col in range(width):
                if col == day_col:
                    continue
                values = [rows[i][col] for i in range(start, end) if rows[i][col]]
                if not values:
                    continue
                header_time = next((TIME_RE.search(rows[r][col]).group() for r in range(start)
                                    if rows[r][col] and TIME_RE.search(rows[r][col])), "")
                if header_time and not any(TIME_RE.search(value) for value in values):
                    values.append(header_time)
                parsed = _parse_course_fragments(values, allow_missing_teacher=True)
                if parsed:
                    subject, teacher, schedule, kind, group, room = parsed
                    result.append(_record(day, subject, teacher, schedule, meta,
                                          source_name, page_no, kind, group, room))
        return result
    return result


def _merge_unique_sessions(sessions):
    unique = OrderedDict()
    for s in sessions:
        key = (s["Jour"], _ascii_key(s["Matière"]), _ascii_key(s["Nom et prénom"]),
               s["Horaire"].replace(" ", "").replace("–", "-"), s["Groupe"])
        unique.setdefault(key, s)
    return list(unique.values())


def _modern_sessions(page, meta, source_name, page_no):
    """Read each drawn timetable lane by position; PyMuPDF tables fragment wrapped cells."""
    lines = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            value = _clean("".join(span["text"] for span in line["spans"]))
            if value:
                x0, y0, x1, y1 = line["bbox"]
                lines.append(((x0 + x1) / 2, (y0 + y1) / 2, value))
    days = sorted([(y, _day(s)) for x, y, s in lines if x < 92 and _day(s)], key=lambda v: v[0])
    if not days:
        return []
    # Positions of the five teaching slots on the landscape A4 template.
    lanes = [(91, 241), (241, 390), (390, 532), (532, 675), (675, 817)]
    out = []
    for index, (day_y, day) in enumerate(days):
        top = (days[index - 1][0] + day_y) / 2 if index else day_y - 20
        bottom = (day_y + days[index + 1][0]) / 2 if index + 1 < len(days) else min(page.rect.height, day_y + 23)
        for left, right in lanes:
            fragments = [s for x, y, s in sorted(lines, key=lambda v: v[1])
                         if left <= x < right and top <= y < bottom]
            time_idx = next((n for n, s in enumerate(fragments) if TIME_RE.search(s)), None)
            if time_idx is None or time_idx == 0:
                continue
            subject = " ".join(s for s in fragments[:time_idx] if s not in ("—", "-"))
            if not subject or _ascii_key(subject) == "EXAMEN":
                continue
            kind_time = fragments[time_idx]
            mt = TIME_RE.search(kind_time)
            if not mt:
                continue
            kind = _clean(kind_time[:mt.start()].strip(" ·"))
            teacher_line = next((s for s in fragments[time_idx + 1:] if "Salle" in s and (TEACHER_RE.match(s) or "·" in s)), "")
            if not teacher_line:
                continue
            parts = re.split(r"\s*·\s*", teacher_line, maxsplit=1)
            teacher = parts[0]
            room = parts[1] if len(parts) > 1 else ""
            out.append(_record(day, subject, teacher, mt.group(), meta, source_name, page_no, kind, room=room))
    return out


def _legacy_sessions(page, meta, source_name, page_no):
    """Read the four-lane CM/TD timetable even if merged day cells shift rows."""
    lines = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            value = _clean("".join(span["text"] for span in line["spans"]))
            if value:
                x0, y0, x1, y1 = line["bbox"]
                lines.append(((x0 + x1) / 2, (y0 + y1) / 2, value))
    days = sorted([(y, _day(s)) for x, y, s in lines if x < 100 and _day(s)])
    if not days:
        return []
    lanes = [(94, 260), (260, 430), (445, 614), (614, 808)]
    out = []
    for index, (day_y, day) in enumerate(days):
        top = (days[index - 1][0] + day_y) / 2 if index else day_y - 36
        bottom = (day_y + days[index + 1][0]) / 2 if index + 1 < len(days) else day_y + 32
        for left, right in lanes:
            fragments = [s for x, y, s in sorted(lines, key=lambda v: v[1])
                         if left <= x < right and top <= y < bottom]
            parsed = _parse_course_fragments(fragments, allow_missing_teacher=True)
            if parsed:
                subject, teacher, schedule, kind, group, room = parsed
                out.append(_record(day, subject, teacher, schedule, meta,
                                   source_name, page_no, kind, group, room))
    return out


def _text_sessions(text, meta, source_name, page_no):
    """Conservative fallback for timetables exported as a list rather than a grid."""
    lines = [_clean(line) for line in text.splitlines() if _clean(line)]
    day_positions = [(i, _day(value)) for i, value in enumerate(lines) if _day(value)]
    out = []
    for pos, (start, day) in enumerate(day_positions):
        stop = day_positions[pos + 1][0] if pos + 1 < len(day_positions) else len(lines)
        chunk = lines[start + 1:stop]
        times = [i for i, value in enumerate(chunk) if TIME_RE.search(value)]
        for n, at in enumerate(times):
            previous = times[n - 1] + 1 if n else 0
            fragments = chunk[max(previous, at - 7):at + 1]
            # Include a teacher and room written immediately below the time.
            fragments.extend(chunk[at + 1:min(len(chunk), at + 3)])
            parsed = _parse_course_fragments(fragments, allow_missing_teacher=True)
            if parsed:
                subject, teacher, schedule, kind, group, room = parsed
                out.append(_record(day, subject, teacher, schedule, meta,
                                   source_name, page_no, kind, group, room))
    return _merge_unique_sessions(out)


def _ocr_page(page, tesseract_cmd=""):
    binary = tesseract_cmd or shutil.which("tesseract")
    if not binary:
        raise RuntimeError("Tesseract OCR non installé ; indiquez le chemin de tesseract.exe.")
    from PIL import Image
    import numpy as np

    scale = 4
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    with tempfile.TemporaryDirectory() as directory:
        path = os.path.join(directory, "emploi.png")
        raw = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
        r, g, b = (raw[:, :, i].astype(np.int16) for i in range(3))
        ink = ((r < 120) & (g < 120) & (b < 120)) | ((r > g + 45) & (g < 150) & (b < 150))
        cleaned_image = Image.fromarray(np.where(ink, 0, 255).astype(np.uint8))
        cleaned_image.save(path)
        for langs in ("fra+eng", "eng"):
            result = subprocess.run([binary, path, "stdout", "-l", langs,
                                     "--psm", "11", "tsv"],
                                    capture_output=True, text=True, timeout=90)
            if result.returncode == 0 and result.stdout.strip():
                grouped = OrderedDict()
                for word in csv.DictReader(io.StringIO(result.stdout), delimiter="\t"):
                    if not word.get("text", "").strip():
                        continue
                    key = tuple(word.get(v, "") for v in ("page_num", "block_num", "par_num", "line_num"))
                    grouped.setdefault(key, []).append(word)
                lines = []
                for words in grouped.values():
                    left = min(int(w["left"]) for w in words)
                    right = max(int(w["left"]) + int(w["width"]) for w in words)
                    top = min(int(w["top"]) for w in words)
                    bottom = max(int(w["top"]) + int(w["height"]) for w in words)
                    lines.append(((left + right) / 2, (top + bottom) / 2,
                                  _clean(" ".join(w["text"] for w in words))))
                days = sorted([(y, value) for x, y, value in lines
                               if x < pix.width * .23 and _day(value)])
                if len(days) < 2:
                    return lines, pix.width
                spacing = min(days[i + 1][0] - days[i][0] for i in range(len(days) - 1))
                rescanned = [(x, y, value) for x, y, value in lines if _day(value)]
                for i, (day_y, _) in enumerate(days):
                    top = ((days[i - 1][0] + day_y) / 2 + .1 * spacing
                           if i else day_y - .36 * spacing)
                    bottom = ((day_y + days[i + 1][0]) / 2 + .1 * spacing
                              if i + 1 < len(days) else day_y + .53 * spacing)
                    xs = sorted(x for x, y, value in lines if top <= y < bottom
                                and x > pix.width * .23 and not _day(value))
                    centers = []
                    for x in xs:
                        if not centers or x - centers[-1][-1] > pix.width * .1:
                            centers.append([x])
                        else:
                            centers[-1].append(x)
                    midpoints = [sum(group) / len(group) for group in centers if len(group) >= 2]
                    for center in midpoints[:8]:
                        distances = [abs(center - other) for other in midpoints if other != center]
                        half = min(min(distances) * .47 if distances else pix.width * .16,
                                   pix.width * .16)
                        x0 = max(0, int(center - half))
                        x1 = min(pix.width, int(center + half))
                        crop = cleaned_image.crop((x0, max(0, int(top)), x1, min(pix.height, int(bottom))))
                        crop_path = os.path.join(directory, "cell.png")
                        crop.save(crop_path)
                        cell_result = subprocess.run([binary, crop_path, "stdout", "-l", langs,
                                                      "--psm", "6"], capture_output=True,
                                                     text=True, timeout=25)
                        cell_lines = [_clean(line) for line in cell_result.stdout.splitlines() if _clean(line)]
                        if cell_result.returncode == 0 and any(TIME_RE.search(s) for s in cell_lines):
                            for j, value in enumerate(cell_lines):
                                rescanned.append((center, top + (j + 1) * (bottom - top) / (len(cell_lines) + 1), value))
                        else:
                            rescanned.extend((x, y, value) for x, y, value in lines
                                             if top <= y < bottom and abs(x - center) < half)
                return rescanned, pix.width
        raise RuntimeError("L'OCR n'a pas pu lire la page (langue ou installation Tesseract).")


def _ocr_spatial_sessions(lines, page_width, meta, source_name, page_no):
    days = sorted([(y, _day(value)) for x, y, value in lines
                   if x < page_width * .23 and _day(value)])
    if not days:
        return []
    day_spacing = min((days[i + 1][0] - days[i][0] for i in range(len(days) - 1)), default=180)
    out = []
    for i, (day_y, day) in enumerate(days):
        top = (days[i - 1][0] + day_y) / 2 if i else day_y - .36 * day_spacing
        bottom = (day_y + days[i + 1][0]) / 2 if i + 1 < len(days) else day_y + .48 * day_spacing
        within = [(x, y, text) for x, y, text in lines if top <= y < bottom]
        anchors = [(x, y) for x, y, text in within if TIME_RE.search(text)]
        for x, y in anchors:
            competitors = [abs(other_x - x) for other_x, _ in anchors if abs(other_x - x) > page_width * .08]
            half_width = min(min(competitors) * .47, page_width * .22) if competitors else page_width * .18
            fragments = [text for col_x, row_y, text in sorted(within, key=lambda item: item[1])
                         if abs(col_x - x) < half_width and col_x > page_width * .23]
            parsed = _parse_course_fragments(fragments, allow_missing_teacher=True)
            if parsed:
                subject, teacher, schedule, kind, group, room = parsed
                out.append(_record(day, subject, teacher, schedule, meta,
                                   source_name, page_no, kind, group, room))
    return _merge_unique_sessions(out)


def extract_sessions_from_pdf_bytes(pdf_bytes, source_name="document.pdf",
                                    ocr_enabled=False, tesseract_cmd=""):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    sessions = []
    warnings = []
    document_meta = extract_metadata(doc[0], source_name) if doc.page_count else {}

    for page_no, page in enumerate(doc, start=1):
        meta = extract_metadata(page, source_name)
        for key in ("filiere", "niveau", "annee"):
            meta[key] = meta[key] or document_meta.get(key, "")
        table = _best_timetable_table(page)
        found = []
        if table is not None:
            data = table.extract()
            if len(data) < 25 and any(_day(row[0] if row else "") for row in data):
                found = _simple_table_sessions(table, meta, source_name, page_no)
            elif any(_clean(c).upper() == "JOUR" for row in data[:6] for c in row if c):
                found = _modern_sessions(page, meta, source_name, page_no)
            elif any(_day(row[0] if row else "") for row in data):
                found = _legacy_sessions(page, meta, source_name, page_no)
        if not found:
            for candidate in page.find_tables().tables:
                generic = _generic_grid_sessions(candidate.extract(), meta, source_name, page_no)
                if len(generic) > len(found):
                    found = generic
            if found:
                warnings.append(f"{source_name} - page {page_no}: grille reconnue automatiquement ; vérifiez les matières et enseignants.")
        raw_text = page.get_text("text")
        if not found and raw_text.strip():
            found = _text_sessions(raw_text, meta, source_name, page_no)
            if found:
                warnings.append(f"{source_name} - page {page_no}: lecture par blocs de texte ; vérifiez chaque séance.")
        if not found and not raw_text.strip() and page.get_images(full=True):
            if ocr_enabled:
                try:
                    ocr_lines, page_width = _ocr_page(page, tesseract_cmd)
                    found = _ocr_spatial_sessions(ocr_lines, page_width, meta,
                                                  source_name, page_no)
                except (OSError, RuntimeError, subprocess.TimeoutExpired) as exc:
                    warnings.append(f"{source_name} - page {page_no}: OCR impossible : {exc}")
                if found:
                    warnings.append(f"{source_name} - page {page_no}: OCR appliqué ; vérifiez tous les champs.")
            else:
                warnings.append(f"{source_name} - page {page_no}: page scannée ; activez l'OCR.")
        sessions.extend(found)
        if not found and (raw_text.strip() or page.get_images(full=True)):
            warnings.append(f"{source_name} - page {page_no}: aucune séance fiable détectée.")

    doc.close()
    return sessions, warnings + _data_warnings(sessions, source_name)


def _data_warnings(sessions, source_name):
    warnings = []
    unknown = [s for s in sessions if s["Nom et prénom"] == "À compléter" or len(_ascii_key(s["Nom et prénom"])) < 5]
    if unknown:
        warnings.append(f"{source_name}: {len(unknown)} séance(s) avec enseignant non renseigné dans l'emploi du temps ; compléter le nom avant export.")
    if any(not s["Filière"] or not s["Niveau"] for s in sessions):
        warnings.append(f"{source_name}: filière ou semestre non reconnu ; vérifier les colonnes Filière et Niveau.")
    return warnings


def extract_sessions_from_docx_bytes(file_bytes, source_name="document.docx"):
    """Read Word tables without requiring LibreOffice on the user's PC."""
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    header = "\n".join(p.text for p in doc.paragraphs)
    class PageText:
        def get_text(self, _):
            return header
    meta = extract_metadata(PageText(), source_name)
    found = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        starts = [i for i, row in enumerate(rows) if row and _day(row[0])]
        table_found = []
        for pos, start in enumerate(starts):
            end = starts[pos + 1] if pos + 1 < len(starts) else len(rows)
            day = _day(rows[start][0])
            ncols = max(map(len, rows[start:end]), default=0)
            for col in range(1, ncols):
                parsed = _parse_course_fragments(row[col] for row in rows[start:end] if col < len(row) and _clean(row[col]))
                if parsed:
                    subject, teacher, schedule, kind, group, room = parsed
                    table_found.append(_record(day, subject, teacher, schedule, meta,
                                               source_name, 1, kind, group, room))
        if not table_found:
            table_found = _generic_grid_sessions(rows, meta, source_name, 1)
        found.extend(table_found)
    used_text_fallback = False
    if not found and header:
        found = _text_sessions(header, meta, source_name, 1)
        used_text_fallback = bool(found)
    unique = OrderedDict()
    for s in found:
        key = (s["Jour"], s["Matière"], s["Nom et prénom"], s["Horaire"], s["Groupe"])
        unique[key] = s
    found = list(unique.values())
    warnings = _data_warnings(found, source_name) if found else [f"{source_name}: aucune séance complète dans les tableaux Word. Vérifiez la mise en page."]
    if used_text_fallback:
        warnings.append(f"{source_name}: texte sans grille reconnu ; vérifiez chaque séance.")
    return found, warnings


def build_intervenants_dataframe(sessions):
    if not sessions:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)

    totals = OrderedDict()
    for session in sessions:
        key = (
            _ascii_key(session.get("Nom et prénom", "")),
            session.get("Filière", ""),
            session.get("Niveau", ""),
            _ascii_key(session.get("Matière", "")),
        )
        if key not in totals:
            totals[key] = {"teacher": session.get("Nom et prénom", ""),
                           "subject": session.get("Matière", ""), "groups": OrderedDict(), "seen": set()}
        duration = session.get("Durée")
        if duration not in (None, ""):
            # Shared sessions for Groups A/B count once for CH, but remain
            # separate in the detailed sheet.
            day_time = (session.get("Jour", ""), session.get("Horaire", ""), session.get("Groupe", ""))
            if day_time in totals[key]["seen"]:
                continue
            totals[key]["seen"].add(day_time)
            group = session.get("Groupe", "")
            try:
                hours = float(str(duration).replace(",", "."))
            except ValueError:
                continue
            totals[key]["groups"][group] = totals[key]["groups"].get(group, 0) + hours

    rows = []
    for (_, filiere, niveau, _), entry in totals.items():
        teacher, subject = entry["teacher"], entry["subject"]
        ch = max(entry["groups"].values(), default=0)
        # Global_Estn stores the nominal weekly load (a 3h15 or 3h30
        # timetable block is recorded as CH=3). Exact durations stay in detail.
        if 2.5 <= ch <= 3.5:
            ch = 3.0
        rows.append({
            "Nom et prénom": teacher,
            "Statut": "",
            "Filière": filiere,
            "Niveau": niveau,
            "Matière": subject,
            "CH": round(ch, 2),
            "Tél": "",
            "Email": "",
            "Département": "",
        })

    return pd.DataFrame(rows, columns=OUTPUT_COLUMNS)


def build_details_dataframe(sessions):
    if not sessions:
        return pd.DataFrame(columns=DETAIL_COLUMNS)
    return pd.DataFrame(sessions).reindex(columns=DETAIL_COLUMNS)


def read_reference_csv(file_bytes):
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            text = file_bytes.decode(encoding)
            return pd.read_csv(io.StringIO(text), sep=None, engine="python", dtype=str).fillna("")
        except Exception as exc:
            last_error = exc
    raise ValueError(f"Impossible de lire le référentiel CSV: {last_error}")


def read_teacher_reference(file_bytes, file_name):
    if file_name.lower().endswith(".xlsx"):
        workbook = pd.ExcelFile(io.BytesIO(file_bytes))
        sheet = "Intervenants" if "Intervenants" in workbook.sheet_names else workbook.sheet_names[0]
        # Global_Estn has a title row and a blank row before its headers.
        preview = pd.read_excel(workbook, sheet_name=sheet, header=None, nrows=5)
        header = next((i for i, row in preview.iterrows() if any(
            _normalize_colname(v) in {"nom et prenom", "enseignant"} for v in row if pd.notna(v)
        )), None)
        if header is None:
            raise ValueError("Colonne « Nom et prénom » absente de la feuille Excel.")
        return pd.read_excel(workbook, sheet_name=sheet, header=header, dtype=str).fillna("")
    return read_reference_csv(file_bytes)


def merge_teacher_reference(intervenants_df, reference_df):
    if intervenants_df.empty or reference_df is None or reference_df.empty:
        return intervenants_df

    aliases = {
        "nom et prenom": "Nom et prénom",
        "nom prenom": "Nom et prénom",
        "enseignant": "Nom et prénom",
        "nom": "Nom et prénom",
        "statut": "Statut",
        "tel": "Tél",
        "telephone": "Tél",
        "mobile": "Tél",
        "email": "Email",
        "mail": "Email",
        "departement": "Département",
        "department": "Département",
    }

    renamed = {}
    for col in reference_df.columns:
        key = _normalize_colname(col)
        if key in aliases:
            renamed[col] = aliases[key]

    ref = reference_df.rename(columns=renamed).copy()
    if "Nom et prénom" not in ref.columns:
        return intervenants_df

    for col in ("Statut", "Tél", "Email", "Département"):
        if col not in ref.columns:
            ref[col] = ""

    ref["_key"] = ref["Nom et prénom"].map(lambda v: " ".join(sorted(_ascii_key(v).split())))
    ref = ref[ref["_key"] != ""]
    ref = ref.drop_duplicates("_key", keep="first")

    out = intervenants_df.copy()
    out["_key"] = out["Nom et prénom"].map(lambda v: " ".join(sorted(_ascii_key(v).split())))

    merged = out.merge(
        ref[["_key", "Statut", "Tél", "Email", "Département"]],
        on="_key",
        how="left",
        suffixes=("", "_ref"),
    )

    for col in ("Statut", "Tél", "Email", "Département"):
        ref_col = f"{col}_ref"
        if ref_col in merged.columns:
            merged[col] = merged[col].where(
                merged[col].astype(str).str.strip() != "",
                merged[ref_col].fillna("")
            )
            merged = merged.drop(columns=[ref_col])

    merged = merged.drop(columns=["_key"])
    return merged.reindex(columns=OUTPUT_COLUMNS).fillna("")




def xlsx_bytes(intervenants_df, details_df=None):
    """Crée un fichier Excel .xlsx en mémoire, prêt pour st.download_button."""
    output = io.BytesIO()

    with xlsxwriter.Workbook(output, {"in_memory": True}) as workbook:
        # ---------- Feuille principale : Intervenants ----------
        ws = workbook.add_worksheet("Intervenants")

        title_fmt = workbook.add_format({
            "bold": True,
            "font_size": 14,
            "align": "left",
            "valign": "vcenter",
        })
        header_fmt = workbook.add_format({
            "bold": True,
            "bg_color": "#D9EAF7",
            "border": 1,
            "align": "center",
            "valign": "vcenter",
            "text_wrap": True,
        })
        text_fmt = workbook.add_format({
            "border": 1,
            "valign": "vcenter",
        })
        ch_fmt = workbook.add_format({
            "border": 1,
            "valign": "vcenter",
            "align": "center",
            "num_format": "0.##",
        })

        columns = list(OUTPUT_COLUMNS)
        ws.merge_range(0, 0, 0, len(columns) - 1, "Liste des Enseignants intervenants", title_fmt)
        ws.set_row(0, 24)
        # Ligne 2 volontairement vide pour reproduire le modèle Global_Estn.

        for col_idx, col_name in enumerate(columns):
            ws.write(2, col_idx, col_name, header_fmt)

        for row_idx, row in enumerate(intervenants_df.reindex(columns=columns).itertuples(index=False, name=None), start=3):
            for col_idx, value in enumerate(row):
                if pd.isna(value):
                    value = ""
                if columns[col_idx] == "CH" and value != "":
                    try:
                        ws.write_number(row_idx, col_idx, float(value), ch_fmt)
                    except Exception:
                        ws.write(row_idx, col_idx, str(value), text_fmt)
                else:
                    ws.write(row_idx, col_idx, value, text_fmt)

        widths = {
            0: 30,  # Nom et prénom
            1: 13,  # Statut
            2: 11,  # Filière
            3: 10,  # Niveau
            4: 40,  # Matière
            5: 8,   # CH
            6: 21,  # Tél
            7: 34,  # Email
            8: 16,  # Département
        }
        for idx, width in widths.items():
            ws.set_column(idx, idx, width)

        last_row = max(3, 2 + len(intervenants_df))
        ws.autofilter(2, 0, last_row, len(columns) - 1)
        ws.freeze_panes(3, 0)
        ws.set_landscape()
        ws.fit_to_pages(1, 0)
        ws.repeat_rows(0, 2)

        # ---------- Deuxième feuille : séances détaillées ----------
        ws2 = workbook.add_worksheet("Séances détaillées")
        details = details_df if details_df is not None else pd.DataFrame(columns=DETAIL_COLUMNS)
        detail_columns = list(DETAIL_COLUMNS)

        for col_idx, col_name in enumerate(detail_columns):
            ws2.write(0, col_idx, col_name, header_fmt)

        for row_idx, row in enumerate(details.reindex(columns=detail_columns).itertuples(index=False, name=None), start=1):
            for col_idx, value in enumerate(row):
                if pd.isna(value):
                    value = ""
                if detail_columns[col_idx] == "Durée" and value != "":
                    try:
                        ws2.write_number(row_idx, col_idx, float(value), ch_fmt)
                    except Exception:
                        ws2.write(row_idx, col_idx, str(value), text_fmt)
                else:
                    ws2.write(row_idx, col_idx, value, text_fmt)

        detail_widths = [13, 38, 10, 29, 17, 10, 24, 24, 12, 10, 19, 28, 8]
        for idx, width in enumerate(detail_widths):
            ws2.set_column(idx, idx, width)
        if len(details) > 0:
            ws2.autofilter(0, 0, len(details), len(detail_columns) - 1)
        ws2.freeze_panes(1, 0)

    output.seek(0)
    return output.getvalue()
