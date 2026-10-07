import pandas as pd
import streamlit as st
import pymupdf as fitz
import io
from edt_changes import changes_dataframe
from ui_navigation import HOME, render_navigation, render_home
from edt_save_state import saved_state, render_saved_state
from edt_correction_state import (
    APPLIED_KEY,
    FLASH_KEY,
    apply_editor_result,
    cancel_editor_changes,
    frame_fingerprint,
    prepare_editor,
    remember_editor_result,
    upload_fingerprint,
)

from app_config import APP_MODE, IS_CLOUD, IS_STREAMLIT_CLOUD
from auth_gate import require_login
from cloud_bootstrap import bootstrap_cloud_database

from agent_edt import ask_agent, answer_factual, ollama_status, OLLAMA_MODEL, AGENT_VERSION
from conflits import detect_conflicts
from rh_ui import render_rh_module
from drive_ui import render_drive_module
from edt_verification_ui import render_global_verification
from edt_statistics_ui import render_edt_statistics
from edt_global_search_ui import render_global_edt_search
from edt_history_ui import render_edt_history
from edt_persistence_ui import sync_edt_changes, render_edt_sync_status
from edt_memory import (
    ensure_edt_memory_db,
    auto_save_uploaded_timetables,
    list_timetables,
    list_versions,
    get_version_id,
    load_version_sessions,
    save_version,
    list_memory,
)

from edt_parser import (
    build_details_dataframe,
    build_intervenants_dataframe,
    extract_sessions_from_pdf_bytes,
    extract_sessions_from_docx_bytes,
    merge_teacher_reference,
    read_teacher_reference,
    xlsx_bytes,
)

st.set_page_config(
    page_title="EDT / RH — Espace de travail",
    page_icon="📄",
    layout="wide",
)

require_login()
_cloud_state = bootstrap_cloud_database()
if IS_CLOUD:
    st.sidebar.caption("☁️ Streamlit Community Cloud" if IS_STREAMLIT_CLOUD else "☁️ Mode Cloud")
    if _cloud_state.get("status") == "error":
        st.error("☁️ Démarrage Cloud interrompu")
        st.error(_cloud_state.get("message"))
        st.info("La base locale temporaire n'est pas utilisée afin d'éviter toute perte ou divergence de données.")
        if st.button("🔄 Réessayer la restauration Google Drive", key="edt_retry_restore"):
            st.rerun()
        st.stop()

try:
    ensure_edt_memory_db()
except Exception as exc:
    st.error(f"Initialisation de la mémoire EDT impossible : {exc}")
    st.stop()

render_edt_sync_status()

module = render_navigation()

with st.sidebar.expander("Fichiers et options d’import", expanded=module == "📅 Emplois du temps"):
    st.header("Documents à importer")
    timetable_files = st.file_uploader(
        "Emplois du temps PDF ou Word",
        type=["pdf", "docx"],
        accept_multiple_files=True,
    )

    reference_file = st.file_uploader(
        "Référentiel enseignants CSV ou Excel (optionnel)",
        type=["csv", "xlsx"],
        help=(
            "Permet de compléter automatiquement Statut, Tél, Email et Département. "
            "Votre Global_Estn.xlsx est accepté (feuille Intervenants), ainsi qu'un CSV "
            "avec Nom et prénom; Statut; Tél; Email; Département. "
            "Seules les identités correspondantes sont complétées."
        ),
    )

    st.header("Lecture des PDF scannés")
    ocr_enabled = st.checkbox("Activer l'OCR si la page est une image", value=False)
    tesseract_cmd = st.text_input(
        "Chemin de tesseract.exe (si nécessaire)", value="",
        help="Exemple Windows : C:\\Program Files\\Tesseract-OCR\\tesseract.exe. "
             "Le programme Tesseract doit être installé séparément."
    ) if ocr_enabled else ""

    st.header("Résultat généré")
    st.caption(
        "Le résultat est généré en fichier Excel .xlsx avec la feuille "
        "« Intervenants » au même format que votre fichier Global_Estn."
    )

if module == HOME:
    render_home()
    st.stop()

if module == "✅ Vérification globale":
    render_global_verification()
    st.stop()

if module == "🔎 Recherche globale":
    render_global_edt_search()
    st.stop()

if module == "📊 Statistiques EDT":
    render_edt_statistics()
    st.stop()

if module == "🗂️ Historique EDT":
    render_edt_history()
    st.stop()

if module == "👥 Ressources humaines":
    render_rh_module()
    st.stop()

if module == "☁️ Google Drive":
    render_drive_module()
    st.stop()

st.title("📄 Emplois du temps PDF et Word → Excel")
st.caption(
    "Importez un ou plusieurs emplois du temps PDF ou Word (.docx). "
    "L'application essaie plusieurs méthodes de lecture, signale les champs incertains "
    "et vous permet de corriger les séances avant de générer le fichier Excel."
)
st.info(
    "Parcours : importez les documents dans la barre latérale, puis suivez les "
    "onglets numérotés pour corriger les séances, vérifier et exporter, puis "
    "enregistrer une version. L'assistant local reste optionnel."
)


if not timetable_files:
    st.subheader("🧠 Emplois du temps mémorisés")

    try:
        ensure_edt_memory_db()
        memory_df = list_memory()

        if memory_df.empty:
            st.info("Aucun emploi du temps n'est encore mémorisé.")
        else:
            st.dataframe(
                memory_df,
                use_container_width=True,
                hide_index=True,
            )

            memory_options = {}

            for _, row in memory_df.iterrows():
                label = (
                    f'{row["Année universitaire"]} | '
                    f'{row["Période"]} | '
                    f'{row["Filière"]} | '
                    f'{row["Niveau"]}'
                )
                memory_options[label] = int(row["id"])

            selected_memory = st.selectbox(
                "Consulter l'historique d'un emploi",
                list(memory_options.keys()),
            )

            timetable_id = memory_options[selected_memory]

            versions_df = list_versions(timetable_id)

            st.markdown("#### Versions disponibles")

            st.dataframe(
                versions_df,
                use_container_width=True,
                hide_index=True,
            )

    except Exception as exc:
        st.warning(f"Mémoire EDT indisponible : {exc}")

    st.divider()

    st.info(
        "Pour importer un nouvel emploi ou créer automatiquement une nouvelle "
        "version, ajoutez un fichier PDF ou Word dans la barre latérale."
    )

    st.stop()

all_sessions = []
all_warnings = []

for uploaded in timetable_files:
    try:
        if uploaded.name.lower().endswith(".docx"):
            sessions, warnings = extract_sessions_from_docx_bytes(uploaded.getvalue(), uploaded.name)
        else:
            sessions, warnings = extract_sessions_from_pdf_bytes(
                uploaded.getvalue(), uploaded.name,
                ocr_enabled=ocr_enabled, tesseract_cmd=tesseract_cmd.strip())
    except Exception as exc:
        sessions, warnings = [], [f"{uploaded.name}: lecture impossible ({exc})."]
    all_sessions.extend(sessions)
    all_warnings.extend(warnings)

if all_warnings:
    with st.expander("⚠️ Vérifications nécessaires avant export", expanded=True):
        for warning in all_warnings:
            st.warning(warning)

if not all_sessions:
    st.warning("Aucune séance n'a été extraite automatiquement. Vous pouvez ajouter les séances "
               "dans l'onglet « 1 · Corriger les séances », ou activer l'OCR pour un PDF scanné.")

with st.expander("Aperçu du document source"):
    if st.checkbox("Afficher l'aperçu", value=False):
        selected_index = st.selectbox("Document", range(len(timetable_files)),
                                      format_func=lambda i: timetable_files[i].name)
        selected = timetable_files[selected_index]
        if selected.name.lower().endswith(".pdf"):
            with fitz.open(stream=selected.getvalue(), filetype="pdf") as document:
                page_index = st.number_input("Page", min_value=1, max_value=len(document), value=1)
                st.image(document[page_index - 1].get_pixmap(matrix=fitz.Matrix(1.6, 1.6)).tobytes("png"),
                         use_column_width=True)
        else:
            from docx import Document
            document = Document(io.BytesIO(selected.getvalue()))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            tables = ["\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
                      for table in document.tables]
            st.text_area("Texte du Word pour vérification", "\n".join(paragraphs + tables), height=300)

metadata = pd.DataFrame([
    {"Document": name,
     "Filière": next((s["Filière"] for s in all_sessions if s["Source PDF"] == name), ""),
     "Niveau": next((s["Niveau"] for s in all_sessions if s["Source PDF"] == name), "")}
    for name in dict.fromkeys(uploaded.name for uploaded in timetable_files)
])
with st.expander("Corriger la filière et le semestre par document", expanded=bool(
        (metadata["Filière"] == "").any() or (metadata["Niveau"] == "").any())):
    edited_metadata = st.data_editor(metadata, hide_index=True, use_container_width=True,
                                     disabled=["Document"], key="document_metadata")
    st.caption("Exemples : Filière = IDSD ; Niveau = S3. Ces corrections s'appliquent à toutes les séances du document.")
metadata_map = edited_metadata.set_index("Document").to_dict("index")
for session in all_sessions:
    values = metadata_map[session["Source PDF"]]
    session["Filière"] = str(values["Filière"]).strip()
    session["Niveau"] = str(values["Niveau"]).strip()


# Mémoire automatique EDT : dès l'upload, après extraction et correction des métadonnées.
# Un contenu strictement identique est reconnu et ne crée pas une nouvelle version.
if all_sessions:
    try:
        _memory_frame = build_details_dataframe(all_sessions)
        _memory_results = auto_save_uploaded_timetables(_memory_frame)

        # === V7.3 CHANGE REPORT ===
        for _memory_item in _memory_results:
            if _memory_item.get('changes_error'):
                st.warning('Version mémorisée, mais comparaison impossible : ' + _memory_item['changes_error'])
            elif _memory_item.get('status') == 'created' and _memory_item.get('changes'):
                _changes = _memory_item['changes']
                if _changes.get('baseline'):
                    st.info(f"🧠 V{_memory_item['version_number']} enregistrée comme version initiale.")
                else:
                    _vn = int(_memory_item['version_number'])
                    st.markdown(f'### 🔄 Changements détectés — V{_vn - 1} → V{_vn}')
                    _c1, _c2, _c3, _c4 = st.columns(4)
                    _c1.metric('Identiques', int(_changes.get('unchanged_count', 0)))
                    _c2.metric('Modifiées', int(_changes.get('modified_count', 0)))
                    _c3.metric('Ajoutées', int(_changes.get('added_count', 0)))
                    _c4.metric('Supprimées', int(_changes.get('removed_count', 0)))
                    _changes_df = changes_dataframe(_changes)
                    if not _changes_df.empty:
                        st.dataframe(_changes_df, width='stretch', hide_index=True)
        _created = [x for x in _memory_results if x.get("status") == "created"]
        _skipped = [x for x in _memory_results if x.get("status") == "skipped"]

        if _created:
            ok_sync, sync_msg = sync_edt_changes()
            for item in _created:
                st.success(
                    f'🧠 {item["filiere"]} {item["niveau"]} — '
                    f'{item["period"]} {item["academic_year"]} : '
                    f'V{item["version_number"]} mémorisée automatiquement.'
                )
            if not ok_sync:
                st.warning(sync_msg)

        for item in _skipped:
            st.warning(
                f'🧠 {item.get("filiere") or "?"} {item.get("niveau") or "?"} : '
                f'{item["message"]}'
            )
    except Exception as exc:
        st.warning(f"🧠 Mémoire automatique EDT non appliquée : {exc}")

correction_tab, export_tab, memory_tab, assistant_tab = st.tabs([
    "1 · Corriger les séances",
    "2 · Vérifier et exporter",
    "3 · Enregistrer / Versions",
    "4 · Assistant local",
])
with correction_tab:
    st.subheader("Séances détectées à vérifier")
    st.caption("Après correction, vérifiez l'état d'export dans l'étape 2. "
               "Corrigez ou supprimez les séances inexactes, et ajoutez celles qui manquent. "
               "Pour une nouvelle ligne, renseignez aussi Jour, Filière, Niveau, Source PDF et Durée. "
               "Si vous modifiez un horaire, corrigez aussi sa durée en heures.")
    baseline_details = build_details_dataframe(all_sessions)
    editor_seed, source_changed = prepare_editor(
        baseline_details, upload_fingerprint(timetable_files), st.session_state
    )
    if source_changed:
        st.info("Les documents importés ont changé : le brouillon de correction a été réinitialisé.")
    edited_details = st.data_editor(editor_seed,
                                    use_container_width=True, hide_index=True,
                                    num_rows="dynamic", key="details_editor")
    remember_editor_result(edited_details, st.session_state)
    applied_details = st.session_state[APPLIED_KEY].copy()
    has_pending_changes = (
        frame_fingerprint(edited_details) != frame_fingerprint(applied_details)
    )
    apply_column, cancel_column = st.columns(2)
    apply_column.button(
        "✅ Appliquer les corrections",
        key="edt_details_apply",
        type="primary",
        disabled=not has_pending_changes,
        use_container_width=True,
        help="Applique le brouillon à l’export, à l’assistant et à l’enregistrement. Ne crée pas encore de version.",
        on_click=apply_editor_result,
        args=(st.session_state, edited_details.copy()),
    )
    cancel_column.button(
        "↩️ Annuler les modifications en cours",
        key="edt_details_reset",
        disabled=not has_pending_changes,
        use_container_width=True,
        help="Revient au dernier brouillon appliqué. Les versions enregistrées ne sont jamais modifiées.",
        on_click=cancel_editor_changes,
        args=(st.session_state,),
    )
    flash = st.session_state.pop(FLASH_KEY, None)
    if flash == "applied":
        st.success("Corrections appliquées au brouillon local. L’export, l’assistant et l’enregistrement utilisent maintenant ces valeurs.")
    elif flash == "cancelled":
        st.info("Modifications en cours annulées. Le dernier brouillon appliqué est restauré.")
    if has_pending_changes:
        st.warning("Modifications en cours non appliquées. Elles sont conservées pendant la navigation, mais l’export et l’enregistrement utilisent encore le dernier brouillon appliqué.")
    else:
        st.caption("Brouillon appliqué localement. Enregistrez une version pour le conserver après fermeture de l’application.")
    correction_save_status = st.empty()
    st.caption("L’état concerne les séances de l’emploi sélectionné dans « Enregistrer / Versions ». "
               "Vérifiez chaque emploi séparément. La provenance Source PDF / Page n’est pas comparée.")

intervenants = build_intervenants_dataframe(
    applied_details.fillna("").to_dict("records"))
incomplete = (applied_details.reindex(columns=["Jour", "Matière", "Nom et prénom", "Horaire",
                                              "Durée", "Filière", "Niveau"])
              .fillna("").astype(str).apply(lambda col: col.str.strip() == "").any(axis=1))
details = applied_details

with export_tab:
    st.subheader("Vérifier la complétude et exporter")
    export_save_status = st.empty()
    st.caption("Télécharger Excel n’enregistre pas une version dans la mémoire locale.")
    incomplete_count = int(incomplete.sum())
    if details.empty:
        st.warning("0 séance(s) à compléter, mais aucune séance n'est disponible pour export.")
    elif incomplete_count:
        st.warning(f"{incomplete_count} séance(s) à compléter avant export.")
    else:
        st.success("Prêt pour export")

    if reference_file is not None:
        try:
            reference = read_teacher_reference(reference_file.getvalue(), reference_file.name)
            intervenants = merge_teacher_reference(intervenants, reference)
            st.success("Référentiel enseignants fusionné.")
        except Exception as exc:
            st.warning(f"Référentiel non appliqué : {exc}")

    c1, c2, c3 = st.columns(3)
    c1.metric("Séances détectées", len(details))
    c2.metric("Lignes enseignants/matières", len(intervenants))
    c3.metric("Documents traités", len(timetable_files))
    if len(set(x["Source PDF"] for x in all_sessions)) < len(timetable_files):
        st.warning("Au moins un document ne contient aucune séance extraite. Vérifiez les avertissements avant l'export.")

    with st.expander("Contrôle par document", expanded=True):
        st.dataframe(details.groupby("Source PDF", as_index=False).agg(
            Séances=("Matière", "size"), Matières=("Matière", "nunique")
        ), use_container_width=True, hide_index=True)
        st.caption("CH reprend la charge hebdomadaire du modèle Global_Estn : "
                   "une séance de 3h15 ou un couple cours/TD de 3h30 est ramené à 3h. "
                   "Les durées exactes figurent dans « Corriger les séances ». "
                   "Vérifiez les enseignants et matières qui ont changé depuis le référentiel.")

    st.subheader("Liste des enseignants intervenants")
    st.caption(
        "Vous pouvez corriger ou compléter les cellules avant de générer le fichier Excel. "
        "Ces modifications du tableau Intervenants concernent uniquement l’export Excel."
    )

    edited = st.data_editor(
        intervenants,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        column_config={
            "CH": st.column_config.NumberColumn(
                "CH",
                min_value=0.0,
                step=0.5,
                format="%.2f",
            )
        },
        key="intervenants_editor",
    )

    excel_data = xlsx_bytes(edited, details)

    st.download_button(
        "⬇️ Télécharger le fichier Excel (.xlsx)",
        data=excel_data,
        file_name="Liste_des_Enseignants_intervenants.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        disabled=intervenants.empty or bool(incomplete.any()),
    )
    st.info(
        "Les séances corrigées sont ajoutées dans une deuxième feuille "
        "du même fichier Excel : « Séances détaillées »."
    )

with assistant_tab:
    st.subheader("🤖 Assistant local Emplois du Temps")
    st.caption(f"Agent local v{AGENT_VERSION} — moteur factuel Python actif")
    if IS_CLOUD:
        st.caption("Mode Cloud : le moteur factuel Python reste disponible. Ollama local sera remplacé par une IA cloud dans une étape suivante.")
    else:
        st.caption(f"Modèle Ollama : {OLLAMA_MODEL}. L'extracteur Python reste la source des données ; l'IA sert à analyser et expliquer.")

    ok, info = ollama_status()
    if ok:
        st.success("Ollama est connecté.")
        if not any(name == OLLAMA_MODEL or name.startswith("r19-mini") for name in info):
            st.warning(f"Le modèle {OLLAMA_MODEL} n'apparaît pas dans Ollama : {', '.join(info) if info else 'aucun modèle'}")
    else:
        st.error(f"Ollama n'est pas joignable : {info}")

    st.markdown("#### Conflits calculés par Python")
    conflicts = detect_conflicts(details)
    if conflicts.empty:
        st.success("Aucun conflit professeur/salle détecté dans les séances actuellement affichées.")
    else:
        st.dataframe(conflicts, use_container_width=True, hide_index=True)

    st.markdown("#### Poser une question à l'assistant")
    question = st.text_area(
        "Question",
        placeholder="Exemples : résume la charge de Pr. X ; quelles salles sont les plus utilisées ? ; explique les conflits détectés.",
        key="ai_question",
        height=100,
    )
    if st.button("🤖 Répondre", use_container_width=True, disabled=not question.strip()):
        try:
            direct = answer_factual(details, question)
            if direct is not None:
                st.info("Source de la réponse : données extraites + moteur factuel Python (Ollama non utilisé).")
                st.markdown(direct)
            elif not ok:
                st.error("Cette question nécessite Ollama, mais Ollama n'est pas joignable.")
            else:
                with st.spinner("Analyse locale avec Ollama..."):
                    answer, truncated = ask_agent(details, question)
                if truncated:
                    st.warning("Le tableau est volumineux : seules les 120 premières séances ont été transmises au modèle pour cette première version.")
                st.info("Source de la réponse : r19-mini à partir des données extraites.")
                st.markdown(answer or "Aucune réponse retournée par le modèle.")
        except Exception as exc:
            st.error(f"Erreur assistant : {exc}")


with memory_tab:
    st.subheader("🧠 Mémoire des emplois du temps")
    st.caption("Chaque nouvel emploi différent devient V1, V2, V3... L'ancienne version reste conservée.")
    st.info(
        "Enregistrer une version conserve les séances corrigées et leur historique. "
        "Cette action ne remplace pas la « Vérification globale », accessible dans la navigation."
    )

    ensure_edt_memory_db()
    mem_details = details.fillna("").copy()

    pairs = []
    if not mem_details.empty:
        pairs = [
            (str(f).strip(), str(n).strip())
            for f, n in mem_details[["Filière", "Niveau"]].drop_duplicates().itertuples(index=False, name=None)
            if str(f).strip() or str(n).strip()
        ]

    if pairs:
        labels = [f"{f or 'Non précisé'} — {n or 'Non précisé'}" for f, n in pairs]
        selected = st.selectbox("Emploi à mémoriser", labels, key="edt_mem_pair")
        filiere, niveau = pairs[labels.index(selected)]
        subset = mem_details[
            (mem_details["Filière"].astype(str).str.strip() == filiere)
            & (mem_details["Niveau"].astype(str).str.strip() == niveau)
        ].copy()

        years = [str(x).strip() for x in subset.get("Année universitaire", pd.Series(dtype=str)).tolist() if str(x).strip()]
        default_year = years[0] if years else "2026-2027"

        c1m, c2m = st.columns(2)
        academic_year = c1m.text_input("Année universitaire", value=default_year, key="edt_mem_year")
        period = c2m.selectbox("Période", ["Automne", "Printemps"], key="edt_mem_period")
        comment = st.text_input("Commentaire de version (optionnel)", key="edt_mem_comment")

        st.caption(f"{len(subset)} séance(s) seront mémorisées pour {filiere} {niveau}.")
        memory_save_status = st.empty()
        if st.button("💾 Enregistrer comme nouvelle version", use_container_width=True, key="edt_mem_save"):
            try:
                result = save_version(subset, academic_year, period, filiere, niveau, comment)
                if result["status"] == "duplicate":
                    st.info(f"Cette version existe déjà : V{result['version_number']}.")
                else:
                    ok_sync, msg_sync = sync_edt_changes()
                    st.success(f"Version V{result['version_number']} enregistrée et définie comme active.")
                    if ok_sync:
                        st.success(msg_sync)
                    else:
                        st.warning(msg_sync)
                if result.get("changes_error"):
                    st.warning("Version enregistrée, mais comparaison impossible : " + result["changes_error"])
            except Exception as exc:
                st.error(f"Enregistrement impossible : {exc}")
        state = saved_state(subset, academic_year, period, filiere, niveau)
        state_label = f"{filiere} {niveau} · {academic_year} · {period}"
        for status_target in (correction_save_status, export_save_status, memory_save_status):
            render_saved_state(status_target, state, state_label)
    else:
        st.info("Aucune Filière / Niveau n'est disponible dans les séances affichées.")
        state = {"status": "empty" if mem_details.empty else "incomplete"}
        for status_target in (correction_save_status, export_save_status):
            render_saved_state(status_target, state, "Séances affichées")

    st.markdown("#### Historique")
    memory = list_memory()
    if memory.empty:
        st.info("Aucun emploi du temps n'est encore mémorisé.")
    else:
        st.dataframe(memory, use_container_width=True, hide_index=True)
        opts = {}
        for _, row in memory.iterrows():
            label = f'{row["Année universitaire"]} | {row["Période"]} | {row["Filière"]} | {row["Niveau"]}'
            opts[label] = int(row["id"])
        chosen = st.selectbox("Voir les versions", list(opts.keys()), key="edt_mem_history")
        st.dataframe(list_versions(opts[chosen]), use_container_width=True, hide_index=True)


st.divider()
st.caption(
    "Le fichier Excel contient deux feuilles : « Intervenants » et « Séances détaillées ». "
    "Statut, Tél, Email et Département peuvent être complétés dans le tableau "
    "ou via un référentiel CSV ou Excel."
)
