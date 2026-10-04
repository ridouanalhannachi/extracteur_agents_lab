import io
import json
import re
import sqlite3
import unicodedata
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, List, Tuple, Optional

import pymupdf as fitz
import pandas as pd
from docx import Document
from PIL import Image

from app_config import DB_PATH

PRIORITY_COLUMNS = [
    "Statut", "Nom", "Prénom", "Email", "Téléphone",
    "Diplôme", "Spécialité", "Département",
]

CANONICAL_TO_DB = {
    "Statut": "statut",
    "Nom": "nom",
    "Prénom": "prenom",
    "Email": "email",
    "Téléphone": "telephone",
    "Diplôme": "diplome",
    "Spécialité": "specialite",
    "Département": "departement",
}

SYNONYMS = {
    "Statut": [
        "statut", "statut enseignant", "type enseignant", "type d enseignant",
        "situation administrative", "qualite", "qualité", "categorie", "catégorie",
        "position administrative", "type intervenant", "nature intervenant",
    ],
    "Nom": [
        "nom", "nom enseignant", "nom de famille", "family name", "surname", "last name",
        "last name required", "familyname", "lastname",
    ],
    "Prénom": [
        "prenom", "prénom", "prenom enseignant", "prénom enseignant", "first name", "given name",
        "first name required", "firstname", "givenname",
    ],
    "Nom complet": [
        "nom et prenom", "nom et prénom", "nom complet", "enseignant", "professeur",
        "intervenant", "full name", "display name", "name",
    ],
    "Email": [
        "email", "e-mail", "e mail", "mail", "adresse mail", "mail professionnel",
        "email professionnel", "adresse electronique", "adresse électronique", "courriel",
        "email address", "email address required", "adresse e-mail", "messagerie",
    ],
    "Téléphone": [
        "telephone", "téléphone", "tel", "tél", "tel portable", "téléphone portable",
        "numero telephone", "numéro téléphone", "gsm", "portable", "mobile", "phone",
        "phone number", "numero de telephone", "numéro de téléphone",
    ],
    "Diplôme": [
        "diplome", "diplôme", "diplome obtenu", "diplôme obtenu", "dernier diplome",
        "dernier diplôme", "plus haut diplome", "plus haut diplôme", "qualification",
        "degree", "highest degree", "grade universitaire", "titre universitaire",
    ],
    "Spécialité": [
        "specialite", "spécialité", "specialite diplome", "spécialité diplôme", "discipline",
        "domaine", "specialization", "spécialisation", "domaine de specialite",
        "domaine de spécialité", "champ disciplinaire",
    ],
    "Département": [
        "departement", "département", "departement d affectation", "département d affectation",
        "dept", "dépt", "department", "structure", "service", "unite", "unité",
        "departement pedagogique", "département pédagogique",
    ],
}


def _norm(value: str) -> str:
    value = unicodedata.normalize("NFD", str(value or ""))
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    value = value.lower().replace("_", " ").replace("-", " ")
    value = re.sub(r"[^a-z0-9@.+ ]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _clean(value) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def _strip_header_annotations(header: str) -> str:
    base = re.sub(r"\[[^\]]*\]", " ", str(header or ""))
    base = re.sub(r"\([^)]*\)", " ", base)
    base = re.sub(r"\b(required|mandatory|obligatoire|read only|readonly)\b", " ", base, flags=re.I)
    return re.sub(r"\s+", " ", base).strip()


def ensure_db(db_path: Path = DB_PATH):
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path) as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS enseignants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                statut TEXT,
                nom TEXT,
                prenom TEXT,
                email TEXT,
                telephone TEXT,
                diplome TEXT,
                specialite TEXT,
                departement TEXT,
                autres_informations TEXT,
                source_document TEXT,
                feuille_section TEXT,
                identity_key TEXT,
                date_import TEXT,
                date_maj TEXT
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_enseignants_email ON enseignants(email)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_enseignants_tel ON enseignants(telephone)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_enseignants_nom ON enseignants(nom, prenom)")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS imports_rh (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_document TEXT,
                nb_lignes INTEGER,
                date_import TEXT
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS fragments_rh (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                donnees_json TEXT,
                source_document TEXT,
                feuille_section TEXT,
                raison TEXT,
                date_import TEXT
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS column_mappings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_header TEXT NOT NULL,
                normalized_header TEXT NOT NULL UNIQUE,
                canonical_field TEXT NOT NULL,
                learned_at TEXT NOT NULL,
                usage_count INTEGER DEFAULT 1
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS unknown_headers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                source_header TEXT NOT NULL,
                normalized_header TEXT NOT NULL UNIQUE,
                source_document TEXT,
                first_seen TEXT NOT NULL,
                last_seen TEXT NOT NULL,
                seen_count INTEGER DEFAULT 1
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS rh_conflicts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                field_name TEXT NOT NULL,
                current_value TEXT,
                incoming_value TEXT,
                source_document TEXT,
                feuille_section TEXT,
                detected_at TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'open',
                resolution TEXT,
                resolved_at TEXT,
                FOREIGN KEY(teacher_id) REFERENCES enseignants(id)
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS conflict_resolution_rules (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                field_name TEXT NOT NULL,
                value_a TEXT,
                value_b TEXT,
                value_a_norm TEXT NOT NULL,
                value_b_norm TEXT NOT NULL,
                preferred_value TEXT,
                preferred_value_norm TEXT,
                decision TEXT NOT NULL,
                created_at TEXT NOT NULL,
                last_used_at TEXT,
                usage_count INTEGER DEFAULT 0,
                UNIQUE(teacher_id, field_name, value_a_norm, value_b_norm),
                FOREIGN KEY(teacher_id) REFERENCES enseignants(id)
            )
        """)
        conn.execute("CREATE INDEX IF NOT EXISTS idx_conflict_rules_teacher ON conflict_resolution_rules(teacher_id, field_name)")
        conn.execute("""
            CREATE TABLE IF NOT EXISTS field_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                teacher_id INTEGER NOT NULL,
                field_name TEXT NOT NULL,
                field_value TEXT,
                source_document TEXT,
                feuille_section TEXT,
                observed_at TEXT NOT NULL,
                FOREIGN KEY(teacher_id) REFERENCES enseignants(id)
            )
        """)
        conn.commit()


def load_column_mappings(db_path: Path = DB_PATH) -> Dict[str, str]:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        rows = conn.execute("SELECT normalized_header, canonical_field FROM column_mappings").fetchall()
    return {r[0]: r[1] for r in rows}


def learn_column_mapping(source_header: str, canonical_field: str, db_path: Path = DB_PATH):
    if canonical_field not in PRIORITY_COLUMNS + ["Nom complet", "Ignorer"]:
        raise ValueError("Champ canonique inconnu")
    ensure_db(db_path)
    nh = _norm(_strip_header_annotations(source_header))
    now = datetime.now().isoformat(timespec="seconds")
    with sqlite3.connect(db_path) as conn:
        conn.execute("""
            INSERT INTO column_mappings(source_header, normalized_header, canonical_field, learned_at, usage_count)
            VALUES (?, ?, ?, ?, 1)
            ON CONFLICT(normalized_header) DO UPDATE SET
                source_header=excluded.source_header,
                canonical_field=excluded.canonical_field,
                learned_at=excluded.learned_at,
                usage_count=column_mappings.usage_count+1
        """, (_clean(source_header), nh, canonical_field, now))
        conn.execute("DELETE FROM unknown_headers WHERE normalized_header=?", (nh,))
        conn.commit()


def register_unknown_header(header: str, source_document: str, db_path: Path = DB_PATH):
    h = _clean(header)
    if not h or h.lower().startswith("unnamed"):
        return
    nh = _norm(_strip_header_annotations(h))
    if not nh:
        return
    ensure_db(db_path)
    now = datetime.now().isoformat(timespec="seconds")
    with sqlite3.connect(db_path) as conn:
        conn.execute("""
            INSERT INTO unknown_headers(source_header, normalized_header, source_document, first_seen, last_seen, seen_count)
            VALUES (?, ?, ?, ?, ?, 1)
            ON CONFLICT(normalized_header) DO UPDATE SET
                source_header=excluded.source_header,
                source_document=excluded.source_document,
                last_seen=excluded.last_seen,
                seen_count=unknown_headers.seen_count+1
        """, (h, nh, source_document, now, now))
        conn.commit()


def load_unknown_headers(db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query("""
            SELECT id, source_header AS 'Colonne source', source_document AS 'Dernière source',
                   seen_count AS 'Occurrences', first_seen AS 'Première vue', last_seen AS 'Dernière vue'
            FROM unknown_headers ORDER BY seen_count DESC, source_header
        """, conn)


def _canonical_header(header: str, learned: Optional[Dict[str, str]] = None):
    original = str(header or "")
    raw_norm = _norm(original)
    base_norm = _norm(_strip_header_annotations(original))

    # Les statuts techniques d'annuaire ne sont pas des statuts RH.
    if raw_norm.startswith("status") and ("read only" in raw_norm or "readonly" in raw_norm):
        return None

    learned = learned if learned is not None else load_column_mappings()
    for candidate in [raw_norm, base_norm]:
        if candidate in learned:
            return None if learned[candidate] == "Ignorer" else learned[candidate]

    candidates = [c for c in [raw_norm, base_norm] if c]
    variants = []
    for canonical, vals in SYNONYMS.items():
        variants.append((canonical, _norm(canonical)))
        variants.extend((canonical, _norm(v)) for v in vals)

    for candidate in candidates:
        for canonical, variant in variants:
            if candidate == variant:
                return canonical

    # Détection floue prudente : utile pour fautes, pluriels et colonnes personnalisées.
    best = (None, 0.0)
    for candidate in candidates:
        if len(candidate) < 3:
            continue
        for canonical, variant in variants:
            ratio = SequenceMatcher(None, candidate, variant).ratio()
            if ratio > best[1]:
                best = (canonical, ratio)
    if best[1] >= 0.88:
        return best[0]
    return None


def _normalize_status(value: str) -> str:
    raw = _clean(value)
    n = _norm(raw)
    if not n:
        return "Non précisé"
    if any(k in n for k in ["vacataire", "vacation", "heure vacataire", "externe"]):
        return "Vacataire"
    if any(k in n for k in ["permanent", "titulaire", "fonctionnaire", "permanente"]):
        return "Permanent"
    return raw


def _split_full_name(value: str) -> Tuple[str, str]:
    raw = _clean(value)
    if not raw:
        return "", ""
    if "," in raw:
        left, right = [x.strip() for x in raw.split(",", 1)]
        return left, right
    parts = raw.split()
    if len(parts) == 2 and parts[0].isupper() and not parts[1].isupper():
        return parts[0], parts[1]
    return raw, ""


def _row_to_record(row: Dict, source: str, sheet: str = "", learned: Optional[Dict[str, str]] = None) -> Dict:
    learned = learned if learned is not None else load_column_mappings()
    mapped = {c: "" for c in PRIORITY_COLUMNS}
    extras = {}
    full_name = ""
    for key, value in row.items():
        key_clean = _clean(key)
        if not key_clean or key_clean.lower().startswith("unnamed"):
            continue
        value_clean = _clean(value)
        if not value_clean:
            continue
        canonical = _canonical_header(key_clean, learned)
        if canonical == "Nom complet":
            full_name = value_clean
        elif canonical in mapped:
            if not mapped[canonical]:
                mapped[canonical] = value_clean
            elif mapped[canonical] != value_clean:
                extras[key_clean] = value_clean
        else:
            extras[key_clean] = value_clean
            register_unknown_header(key_clean, source)

    if full_name and not mapped["Nom"] and not mapped["Prénom"]:
        mapped["Nom"], mapped["Prénom"] = _split_full_name(full_name)
        extras.setdefault("Nom complet source", full_name)
    elif full_name:
        extras.setdefault("Nom complet source", full_name)

    mapped["Statut"] = _normalize_status(mapped["Statut"])
    mapped["Autres informations"] = json.dumps(extras, ensure_ascii=False, sort_keys=True) if extras else "{}"
    mapped["Source document"] = source
    mapped["Feuille/Section"] = sheet
    return mapped


def _record_has_identity(rec: Dict) -> bool:
    return any(_clean(rec.get(c)) for c in ["Nom", "Prénom", "Email", "Téléphone"])


def _record_has_any_rh(rec: Dict) -> bool:
    for c in PRIORITY_COLUMNS:
        value = _clean(rec.get(c))
        if value and value != "Non précisé":
            return True
    try:
        extras = json.loads(rec.get("Autres informations") or "{}")
        return bool(extras)
    except Exception:
        return bool(_clean(rec.get("Autres informations")))


def _score_header_row(values: List[str], learned=None) -> int:
    return sum(1 for value in values if _canonical_header(value, learned))


def _decode_text_bytes(raw: bytes):
    for encoding in ["utf-8-sig", "utf-8", "cp1252", "latin1"]:
        try:
            return raw.decode(encoding), encoding
        except Exception:
            continue
    return None, None


def _read_csv_flexible(text: str, learned=None):
    candidates = []
    for sep in [None, ";", ",", "\t", "|"]:
        try:
            kwargs = {"header": None, "dtype": str}
            if sep is None:
                kwargs.update({"sep": None, "engine": "python"})
            else:
                kwargs.update({"sep": sep})
            df = pd.read_csv(io.StringIO(text), **kwargs)
            if df is not None and not df.empty:
                candidates.append(df)
        except Exception:
            pass
    if not candidates:
        return None, -1, 0
    best_df, best_idx, best_score = candidates[0], 0, -1
    for df in candidates:
        for idx in range(min(30, len(df))):
            vals = [_clean(v) for v in df.iloc[idx].tolist()]
            score = _score_header_row(vals, learned)
            if score > best_score:
                best_df, best_idx, best_score = df, idx, score
    return best_df, best_idx, best_score


def _read_excel_sheet(raw: bytes, sheet_name, learned=None, engine=None) -> pd.DataFrame:
    preview = pd.read_excel(io.BytesIO(raw), sheet_name=sheet_name, header=None, engine=engine)
    best_idx, best_score = 0, -1
    for idx in range(min(20, len(preview))):
        score = _score_header_row([_clean(x) for x in preview.iloc[idx].tolist()], learned)
        if score > best_score:
            best_idx, best_score = idx, score
    return pd.read_excel(io.BytesIO(raw), sheet_name=sheet_name, header=best_idx, engine=engine)


def extract_from_excel(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    warnings, records = [], []
    learned = load_column_mappings()
    try:
        xls = pd.ExcelFile(io.BytesIO(raw))
        for sheet in xls.sheet_names:
            try:
                df = _read_excel_sheet(raw, sheet, learned=learned)
                df = df.dropna(how="all")
                for _, row in df.iterrows():
                    rec = _row_to_record(row.to_dict(), filename, str(sheet), learned)
                    if _record_has_any_rh(rec):
                        records.append(rec)
            except Exception as exc:
                warnings.append(f"{filename} / feuille {sheet}: lecture partielle impossible ({exc}).")
    except Exception as exc:
        warnings.append(f"{filename}: lecture Excel impossible ({exc}).")
    return records, warnings


def extract_from_csv(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    text, encoding = _decode_text_bytes(raw)
    if text is None:
        return [], [f"{filename}: encodage texte non reconnu."]
    learned = load_column_mappings()
    raw_df, header_idx, header_score = _read_csv_flexible(text, learned)
    if raw_df is None:
        return [], [f"{filename}: lecture CSV/TSV impossible."]
    if header_score <= 0:
        sample = [str(x) for x in raw_df.iloc[0].tolist()[:12]] if len(raw_df) else []
        for h in sample:
            register_unknown_header(h, filename)
        return [], [
            f"{filename}: aucune colonne RH reconnue automatiquement. En-tête aperçu : {sample}. "
            "Vous pouvez apprendre ces colonnes dans « Mémoire des colonnes »."
        ]
    headers = [_clean(x) or f"Colonne_{i+1}" for i, x in enumerate(raw_df.iloc[header_idx].tolist())]
    data = raw_df.iloc[header_idx + 1:].copy()
    data.columns = headers
    data = data.dropna(how="all")
    records = []
    for _, row in data.iterrows():
        rec = _row_to_record(row.to_dict(), filename, f"CSV/TSV (encodage {encoding}, en-tête ligne {header_idx + 1})", learned)
        if _record_has_any_rh(rec):
            records.append(rec)
    warnings = []
    if header_idx > 0:
        warnings.append(f"{filename}: en-tête RH détecté automatiquement à la ligne {header_idx + 1}.")
    return records, warnings


def extract_from_json(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    try:
        data = json.loads(raw.decode("utf-8-sig"))
    except Exception as exc:
        return [], [f"{filename}: JSON illisible ({exc})."]
    if isinstance(data, dict):
        for key in ["data", "users", "enseignants", "teachers", "rows", "items"]:
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    if not isinstance(data, list):
        return [], [f"{filename}: structure JSON non reconnue."]
    learned = load_column_mappings()
    records = []
    for item in data:
        if isinstance(item, dict):
            rec = _row_to_record(item, filename, "JSON", learned)
            if _record_has_any_rh(rec):
                records.append(rec)
    return records, [] if records else [f"{filename}: aucune fiche RH reconnue dans le JSON."]


def _pairs_from_text(text: str) -> Dict:
    pairs = {}
    lines = [re.sub(r"\s+", " ", x).strip() for x in text.splitlines() if x.strip()]
    for line in lines:
        m = re.match(r"^([^:;=]{2,80})\s*[:=]\s*(.+)$", line)
        if m:
            pairs[m.group(1).strip()] = m.group(2).strip()
    return pairs


def _records_from_text_blocks(text: str, filename: str, section_prefix="Texte") -> List[Dict]:
    learned = load_column_mappings()
    blocks = [b.strip() for b in re.split(r"\n\s*\n+", text) if b.strip()]
    if not blocks:
        blocks = [text]
    records = []
    for i, block in enumerate(blocks, 1):
        pairs = _pairs_from_text(block)
        if pairs:
            rec = _row_to_record(pairs, filename, f"{section_prefix} {i}", learned)
            if _record_has_any_rh(rec):
                records.append(rec)
    if records:
        return records
    # Fallback : document entier sous forme clé/valeur.
    pairs = _pairs_from_text(text)
    if pairs:
        rec = _row_to_record(pairs, filename, section_prefix, learned)
        if _record_has_any_rh(rec):
            records.append(rec)
    return records


def _records_from_docx_tables(document: Document, filename: str) -> List[Dict]:
    learned = load_column_mappings()
    records = []
    for t_idx, table in enumerate(document.tables, start=1):
        rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        scores = [_score_header_row(row, learned) for row in rows[:10]]
        header_idx = max(range(len(scores)), key=lambda i: scores[i]) if scores else 0
        if scores and scores[header_idx] >= 1:
            headers = rows[header_idx]
            for row in rows[header_idx + 1:]:
                data = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
                rec = _row_to_record(data, filename, f"Tableau {t_idx}", learned)
                if _record_has_any_rh(rec):
                    records.append(rec)
        elif all(len(row) >= 2 for row in rows) and sum(1 for row in rows if _canonical_header(row[0], learned)) >= 1:
            data = {row[0]: row[1] for row in rows if len(row) >= 2}
            rec = _row_to_record(data, filename, f"Tableau {t_idx}", learned)
            if _record_has_any_rh(rec):
                records.append(rec)
    return records


def extract_from_docx(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    try:
        doc = Document(io.BytesIO(raw))
    except Exception as exc:
        return [], [f"{filename}: lecture Word impossible ({exc})."]
    records = _records_from_docx_tables(doc, filename)
    paragraphs = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    records.extend(_records_from_text_blocks(paragraphs, filename, "Texte Word"))
    records = deduplicate_preview(pd.DataFrame(records)).to_dict("records") if records else []
    return records, [] if records else [f"{filename}: aucun enregistrement RH structuré détecté automatiquement."]


def _ocr_image(image: Image.Image) -> Tuple[str, Optional[str]]:
    try:
        import pytesseract
        return pytesseract.image_to_string(image, lang="fra+eng"), None
    except Exception as exc:
        return "", str(exc)


def extract_from_image(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    try:
        image = Image.open(io.BytesIO(raw)).convert("RGB")
    except Exception as exc:
        return [], [f"{filename}: image illisible ({exc})."]
    text, err = _ocr_image(image)
    if err:
        return [], [
            f"{filename}: OCR indisponible ({err}). Installez Tesseract : "
            "sudo apt install tesseract-ocr tesseract-ocr-fra"
        ]
    records = _records_from_text_blocks(text, filename, "OCR image")
    return records, [] if records else [f"{filename}: OCR effectué mais aucune fiche RH reconnue automatiquement."]


def extract_from_pdf(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    records, warnings = [], []
    learned = load_column_mappings()
    try:
        with fitz.open(stream=raw, filetype="pdf") as doc:
            for p_idx, page in enumerate(doc, start=1):
                page_text = page.get_text("text") or ""
                records.extend(_records_from_text_blocks(page_text, filename, f"Page {p_idx}"))
                try:
                    finder = page.find_tables()
                    for t_idx, table in enumerate(getattr(finder, "tables", []), start=1):
                        data = table.extract()
                        if not data:
                            continue
                        rows = [[_clean(c) for c in row] for row in data]
                        scores = [_score_header_row(row, learned) for row in rows[:10]]
                        header_idx = max(range(len(scores)), key=lambda i: scores[i]) if scores else 0
                        if scores and scores[header_idx] >= 1:
                            headers = rows[header_idx]
                            for row in rows[header_idx + 1:]:
                                rec = _row_to_record(
                                    {headers[i]: row[i] for i in range(min(len(headers), len(row)))},
                                    filename, f"Page {p_idx} tableau {t_idx}", learned,
                                )
                                if _record_has_any_rh(rec):
                                    records.append(rec)
                except Exception:
                    pass
                if len(page_text.strip()) < 40:
                    try:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False)
                        image = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text, err = _ocr_image(image)
                        if not err:
                            records.extend(_records_from_text_blocks(ocr_text, filename, f"Page {p_idx} OCR"))
                    except Exception:
                        pass
    except Exception as exc:
        warnings.append(f"{filename}: lecture PDF impossible ({exc}).")
    if records:
        records = deduplicate_preview(pd.DataFrame(records)).to_dict("records")
    elif not warnings:
        warnings.append(f"{filename}: aucune fiche RH reconnue. Si le PDF est scanné, installez Tesseract OCR.")
    return records, warnings


def extract_from_txt(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    text, _ = _decode_text_bytes(raw)
    if text is None:
        return [], [f"{filename}: texte illisible."]
    records = _records_from_text_blocks(text, filename, "Texte")
    return records, [] if records else [f"{filename}: aucune fiche RH structurée détectée."]


def extract_rh_file(raw: bytes, filename: str) -> Tuple[List[Dict], List[str]]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm", ".xls", ".ods")):
        return extract_from_excel(raw, filename)
    if lower.endswith((".csv", ".tsv")):
        return extract_from_csv(raw, filename)
    if lower.endswith(".json"):
        return extract_from_json(raw, filename)
    if lower.endswith(".docx"):
        return extract_from_docx(raw, filename)
    if lower.endswith(".pdf"):
        return extract_from_pdf(raw, filename)
    if lower.endswith((".txt", ".md")):
        return extract_from_txt(raw, filename)
    if lower.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")):
        return extract_from_image(raw, filename)
    return [], [f"{filename}: format non pris en charge actuellement."]


def preview_tabular_file(raw: bytes, filename: str) -> pd.DataFrame:
    lower = filename.lower()
    try:
        if lower.endswith((".csv", ".tsv")):
            text, _ = _decode_text_bytes(raw)
            if text is None:
                return pd.DataFrame()
            df, _, _ = _read_csv_flexible(text, load_column_mappings())
            return df.head(30).fillna("") if df is not None else pd.DataFrame()
        if lower.endswith((".xlsx", ".xls", ".xlsm", ".ods")):
            return pd.read_excel(io.BytesIO(raw), header=None).head(30).fillna("")
    except Exception:
        return pd.DataFrame()
    return pd.DataFrame()


def _identity_key(row: Dict) -> str:
    email = _norm(row.get("Email", ""))
    if email:
        return f"email:{email}"
    phone = re.sub(r"\D+", "", str(row.get("Téléphone", "")))
    if len(phone) >= 8:
        return f"tel:{phone}"
    name = _norm(f"{row.get('Nom', '')} {row.get('Prénom', '')}")
    return f"nom:{name}" if name else ""


def _same_person(a: Dict, b: Dict) -> bool:
    ea, eb = _norm(a.get("Email", "")), _norm(b.get("Email", ""))
    if ea and eb and ea == eb:
        return True
    pa = re.sub(r"\D+", "", str(a.get("Téléphone", "")))
    pb = re.sub(r"\D+", "", str(b.get("Téléphone", "")))
    if len(pa) >= 8 and len(pb) >= 8 and pa == pb:
        return True
    na = _norm(f"{a.get('Nom', '')} {a.get('Prénom', '')}")
    nb = _norm(f"{b.get('Nom', '')} {b.get('Prénom', '')}")
    return bool(na and nb and na == nb)


def _merge_json_text(a: str, b: str) -> str:
    try:
        da = json.loads(a or "{}") if isinstance(a, str) else {}
    except Exception:
        da = {"Ancienne valeur": _clean(a)} if _clean(a) else {}
    try:
        db = json.loads(b or "{}") if isinstance(b, str) else {}
    except Exception:
        db = {"Nouvelle valeur": _clean(b)} if _clean(b) else {}
    if not isinstance(da, dict):
        da = {"Ancienne valeur": da}
    if not isinstance(db, dict):
        db = {"Nouvelle valeur": db}
    for k, v in db.items():
        if v in [None, ""]:
            continue
        if k not in da or da[k] in [None, ""]:
            da[k] = v
        elif da[k] != v:
            existing = da[k] if isinstance(da[k], list) else [da[k]]
            if v not in existing:
                existing.append(v)
            da[k] = existing
    return json.dumps(da, ensure_ascii=False, sort_keys=True)


def _merge_sources(a: str, b: str) -> str:
    vals = []
    for part in [a, b]:
        for x in str(part or "").split(" ; "):
            x = x.strip()
            if x and x not in vals:
                vals.append(x)
    return " ; ".join(vals)


def _merge_record(current: Dict, incoming: Dict) -> Dict:
    out = dict(current)
    for col, value in incoming.items():
        if col == "Autres informations":
            out[col] = _merge_json_text(out.get(col, "{}"), value or "{}")
        elif col in ["Source document", "Feuille/Section"]:
            out[col] = _merge_sources(out.get(col, ""), value)
        elif (not _clean(out.get(col)) or out.get(col) == "Non précisé") and _clean(value) and value != "Non précisé":
            out[col] = value
    return out


def deduplicate_preview(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    rows = []
    for _, row in df.fillna("").iterrows():
        data = row.to_dict()
        matched = None
        if _record_has_identity(data):
            for i, current in enumerate(rows):
                if _same_person(current, data):
                    matched = i
                    break
        if matched is None:
            rows.append(data)
        else:
            rows[matched] = _merge_record(rows[matched], data)
    return pd.DataFrame(rows)


def _find_existing_teacher(conn, data: Dict):
    rows = conn.execute("""
        SELECT id, statut, nom, prenom, email, telephone, diplome, specialite,
               departement, autres_informations, source_document, feuille_section, identity_key
        FROM enseignants
    """).fetchall()
    incoming = {"Nom": data.get("Nom", ""), "Prénom": data.get("Prénom", ""),
                "Email": data.get("Email", ""), "Téléphone": data.get("Téléphone", "")}
    for r in rows:
        existing = {"Nom": r[2] or "", "Prénom": r[3] or "", "Email": r[4] or "", "Téléphone": r[5] or ""}
        if _same_person(existing, incoming):
            return r
    return None


def _add_history(conn, teacher_id: int, field_name: str, value: str, source: str, section: str, now: str):
    if not _clean(value) or value == "Non précisé":
        return
    conn.execute("""
        INSERT INTO field_history(teacher_id, field_name, field_value, source_document, feuille_section, observed_at)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (teacher_id, field_name, _clean(value), _clean(source), _clean(section), now))


def _conflict_pair(current: str, incoming: str):
    pairs = [(_norm(current), _clean(current)), (_norm(incoming), _clean(incoming))]
    pairs.sort(key=lambda x: x[0])
    return pairs[0][0], pairs[1][0], pairs[0][1], pairs[1][1]


def _remember_conflict_rule(conn, teacher_id: int, field: str, current: str, incoming: str,
                            preferred_value: str, decision: str, now: str):
    a_norm, b_norm, a_val, b_val = _conflict_pair(current, incoming)
    conn.execute("""
        INSERT INTO conflict_resolution_rules(
            teacher_id, field_name, value_a, value_b, value_a_norm, value_b_norm,
            preferred_value, preferred_value_norm, decision, created_at, last_used_at, usage_count
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 0)
        ON CONFLICT(teacher_id, field_name, value_a_norm, value_b_norm) DO UPDATE SET
            value_a=excluded.value_a,
            value_b=excluded.value_b,
            preferred_value=excluded.preferred_value,
            preferred_value_norm=excluded.preferred_value_norm,
            decision=excluded.decision,
            created_at=excluded.created_at
    """, (
        int(teacher_id), field, a_val, b_val, a_norm, b_norm,
        _clean(preferred_value), _norm(preferred_value), decision, now,
    ))


def _find_conflict_rule(conn, teacher_id: int, field: str, current: str, incoming: str):
    a_norm, b_norm, _, _ = _conflict_pair(current, incoming)
    return conn.execute("""
        SELECT id, preferred_value, decision
        FROM conflict_resolution_rules
        WHERE teacher_id=? AND field_name=? AND value_a_norm=? AND value_b_norm=?
    """, (int(teacher_id), field, a_norm, b_norm)).fetchone()


def _mark_conflict_rule_used(conn, rule_id: int, now: str):
    conn.execute("""
        UPDATE conflict_resolution_rules
        SET usage_count=usage_count+1, last_used_at=?
        WHERE id=?
    """, (now, int(rule_id)))


def _record_auto_resolution(conn, teacher_id: int, field: str, current: str, incoming: str,
                            source: str, section: str, decision: str, preferred: str, now: str):
    label = {
        "actuelle": "Résolu automatiquement par règle mémorisée : valeur préférée conservée",
        "nouvelle": "Résolu automatiquement par règle mémorisée : valeur préférée adoptée",
        "les_deux": "Résolu automatiquement par règle mémorisée : valeur préférée + alternative archivée",
    }.get(decision, "Résolu automatiquement par règle mémorisée")
    conn.execute("""
        INSERT INTO rh_conflicts(
            teacher_id, field_name, current_value, incoming_value, source_document,
            feuille_section, detected_at, status, resolution, resolved_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, 'auto_resolved', ?, ?)
    """, (
        int(teacher_id), field, _clean(current), _clean(incoming), _clean(source), _clean(section),
        now, f"{label} — valeur retenue : {_clean(preferred)}", now,
    ))


def _open_conflict(conn, teacher_id: int, field: str, current: str, incoming: str, source: str, section: str, now: str) -> bool:
    if _norm(current) == _norm(incoming):
        return False
    exists = conn.execute("""
        SELECT id FROM rh_conflicts
        WHERE teacher_id=? AND field_name=? AND incoming_value=? AND status='open'
    """, (teacher_id, field, _clean(incoming))).fetchone()
    if exists:
        return False
    conn.execute("""
        INSERT INTO rh_conflicts(teacher_id, field_name, current_value, incoming_value,
                                 source_document, feuille_section, detected_at, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'open')
    """, (teacher_id, field, _clean(current), _clean(incoming), _clean(source), _clean(section), now))
    return True


def save_teachers(df: pd.DataFrame, db_path: Path = DB_PATH) -> Tuple[int, int, int, int]:
    ensure_db(db_path)
    inserted = updated = pending = conflicts_created = 0
    now = datetime.now().isoformat(timespec="seconds")
    with sqlite3.connect(db_path) as conn:
        for _, row in df.fillna("").iterrows():
            data = row.to_dict()
            if not _record_has_any_rh(data):
                continue
            key = _identity_key(data)
            if not key:
                conn.execute("""
                    INSERT INTO fragments_rh(donnees_json, source_document, feuille_section, raison, date_import)
                    VALUES (?, ?, ?, ?, ?)
                """, (json.dumps(data, ensure_ascii=False), _clean(data.get("Source document")),
                      _clean(data.get("Feuille/Section")), "Aucun identifiant enseignant", now))
                pending += 1
                continue

            existing = _find_existing_teacher(conn, data)
            if existing:
                teacher_id = int(existing[0])
                existing_data = {
                    "Statut": existing[1] or "", "Nom": existing[2] or "", "Prénom": existing[3] or "",
                    "Email": existing[4] or "", "Téléphone": existing[5] or "", "Diplôme": existing[6] or "",
                    "Spécialité": existing[7] or "", "Département": existing[8] or "",
                    "Autres informations": existing[9] or "{}", "Source document": existing[10] or "",
                    "Feuille/Section": existing[11] or "",
                }
                merged = dict(existing_data)
                changed = False
                for field in PRIORITY_COLUMNS:
                    incoming = _clean(data.get(field))
                    current = _clean(existing_data.get(field))
                    if not incoming or incoming == "Non précisé":
                        continue
                    _add_history(conn, teacher_id, field, incoming, data.get("Source document", ""), data.get("Feuille/Section", ""), now)
                    if not current or current == "Non précisé":
                        merged[field] = incoming
                        changed = True
                    elif _norm(current) != _norm(incoming):
                        rule = _find_conflict_rule(conn, teacher_id, field, current, incoming)
                        if rule:
                            rule_id, preferred, remembered_decision = rule
                            preferred = _clean(preferred) or current
                            if remembered_decision == "les_deux":
                                merged[field] = preferred
                                alternative = incoming if _norm(preferred) == _norm(current) else current
                                alt = json.dumps({f"Alternative {field}": alternative}, ensure_ascii=False)
                                merged["Autres informations"] = _merge_json_text(
                                    merged.get("Autres informations", "{}"), alt
                                )
                            else:
                                merged[field] = preferred
                            changed = True
                            _mark_conflict_rule_used(conn, rule_id, now)
                            _record_auto_resolution(
                                conn, teacher_id, field, current, incoming,
                                data.get("Source document", ""), data.get("Feuille/Section", ""),
                                remembered_decision, preferred, now,
                            )
                        elif _open_conflict(conn, teacher_id, field, current, incoming,
                                            data.get("Source document", ""), data.get("Feuille/Section", ""), now):
                            conflicts_created += 1
                merged["Autres informations"] = _merge_json_text(merged.get("Autres informations", "{}"), data.get("Autres informations", "{}"))
                merged["Source document"] = _merge_sources(existing_data.get("Source document", ""), data.get("Source document", ""))
                merged["Feuille/Section"] = _merge_sources(existing_data.get("Feuille/Section", ""), data.get("Feuille/Section", ""))
                conn.execute("""
                    UPDATE enseignants SET statut=?, nom=?, prenom=?, email=?, telephone=?, diplome=?, specialite=?,
                        departement=?, autres_informations=?, source_document=?, feuille_section=?, date_maj=?
                    WHERE id=?
                """, (
                    _clean(merged.get("Statut")) or "Non précisé", _clean(merged.get("Nom")), _clean(merged.get("Prénom")),
                    _clean(merged.get("Email")), _clean(merged.get("Téléphone")), _clean(merged.get("Diplôme")),
                    _clean(merged.get("Spécialité")), _clean(merged.get("Département")),
                    _clean(merged.get("Autres informations")) or "{}", _clean(merged.get("Source document")),
                    _clean(merged.get("Feuille/Section")), now, teacher_id,
                ))
                updated += 1
            else:
                cur = conn.execute("""
                    INSERT INTO enseignants (
                        statut, nom, prenom, email, telephone, diplome, specialite, departement,
                        autres_informations, source_document, feuille_section, identity_key, date_import, date_maj
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    _clean(data.get("Statut")) or "Non précisé", _clean(data.get("Nom")), _clean(data.get("Prénom")),
                    _clean(data.get("Email")), _clean(data.get("Téléphone")), _clean(data.get("Diplôme")),
                    _clean(data.get("Spécialité")), _clean(data.get("Département")), _clean(data.get("Autres informations")) or "{}",
                    _clean(data.get("Source document")), _clean(data.get("Feuille/Section")), key, now, now,
                ))
                teacher_id = int(cur.lastrowid)
                for field in PRIORITY_COLUMNS:
                    _add_history(conn, teacher_id, field, data.get(field, ""), data.get("Source document", ""), data.get("Feuille/Section", ""), now)
                inserted += 1
        if "Source document" in df.columns:
            for source, group in df.groupby("Source document", dropna=False):
                conn.execute("INSERT INTO imports_rh(source_document, nb_lignes, date_import) VALUES (?, ?, ?)",
                             (_clean(source), int(len(group)), now))
        conn.commit()
    return inserted, updated, pending, conflicts_created


def load_conflicts(status: str = "open", db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    where = "WHERE c.status=?" if status != "all" else ""
    params = (status,) if status != "all" else ()
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query(f"""
            SELECT c.id, c.teacher_id AS 'ID enseignant',
                   TRIM(COALESCE(e.nom,'') || ' ' || COALESCE(e.prenom,'')) AS Enseignant,
                   c.field_name AS Champ, c.current_value AS 'Valeur actuelle',
                   c.incoming_value AS 'Nouvelle valeur', c.source_document AS Source,
                   c.feuille_section AS Section, c.detected_at AS 'Détecté le',
                   c.status AS Statut, c.resolution AS Résolution
            FROM rh_conflicts c JOIN enseignants e ON e.id=c.teacher_id
            {where}
            ORDER BY c.id DESC
        """, conn, params=params)


def resolve_conflict(conflict_id: int, decision: str, db_path: Path = DB_PATH):
    ensure_db(db_path)
    now = datetime.now().isoformat(timespec="seconds")
    with sqlite3.connect(db_path) as conn:
        row = conn.execute("""
            SELECT teacher_id, field_name, current_value, incoming_value, status
            FROM rh_conflicts WHERE id=?
        """, (int(conflict_id),)).fetchone()
        if not row or row[4] != "open":
            return
        teacher_id, field, current, incoming, _ = row
        preferred = current
        if decision == "nouvelle":
            preferred = incoming
            db_col = CANONICAL_TO_DB.get(field)
            if db_col:
                conn.execute(f"UPDATE enseignants SET {db_col}=?, date_maj=? WHERE id=?", (_clean(incoming), now, teacher_id))
        elif decision == "les_deux":
            preferred = current
            teacher = conn.execute("SELECT autres_informations FROM enseignants WHERE id=?", (teacher_id,)).fetchone()
            extra = teacher[0] if teacher else "{}"
            alt = json.dumps({f"Alternative {field}": incoming}, ensure_ascii=False)
            merged = _merge_json_text(extra, alt)
            conn.execute("UPDATE enseignants SET autres_informations=?, date_maj=? WHERE id=?", (merged, now, teacher_id))
        elif decision == "actuelle":
            preferred = current

        # Apprentissage permanent : même enseignant + même champ + même paire de valeurs.
        # Si le même conflit réapparait (même avec valeurs inversées), il sera résolu automatiquement.
        _remember_conflict_rule(conn, teacher_id, field, current, incoming, preferred, decision, now)

        resolution = {
            "actuelle": "Valeur actuelle conservée — décision mémorisée",
            "nouvelle": "Nouvelle valeur adoptée — décision mémorisée",
            "les_deux": "Actuelle conservée + alternative archivée — décision mémorisée",
        }.get(decision, decision)
        conn.execute("UPDATE rh_conflicts SET status='resolved', resolution=?, resolved_at=? WHERE id=?",
                     (resolution, now, int(conflict_id)))
        conn.commit()


def load_conflict_rules(db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query("""
            SELECT r.id,
                   TRIM(COALESCE(e.nom,'') || ' ' || COALESCE(e.prenom,'')) AS Enseignant,
                   r.field_name AS Champ,
                   r.value_a AS 'Valeur A',
                   r.value_b AS 'Valeur B',
                   r.preferred_value AS 'Valeur préférée',
                   CASE r.decision
                       WHEN 'actuelle' THEN 'Conserver la valeur préférée'
                       WHEN 'nouvelle' THEN 'Adopter la valeur préférée'
                       WHEN 'les_deux' THEN 'Valeur préférée + archiver alternative'
                       ELSE r.decision
                   END AS Décision,
                   r.usage_count AS 'Applications auto',
                   r.created_at AS 'Mémorisée le',
                   r.last_used_at AS 'Dernière application'
            FROM conflict_resolution_rules r
            JOIN enseignants e ON e.id=r.teacher_id
            ORDER BY r.id DESC
        """, conn)


def delete_conflict_rule(rule_id: int, db_path: Path = DB_PATH):
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        conn.execute("DELETE FROM conflict_resolution_rules WHERE id=?", (int(rule_id),))
        conn.commit()


def load_fragments(db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query("""
            SELECT id, donnees_json AS 'Données', source_document AS 'Source document',
                   feuille_section AS 'Feuille/Section', raison AS 'Raison', date_import AS 'Date import'
            FROM fragments_rh ORDER BY id DESC
        """, conn)


def load_teachers(db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query("""
            SELECT id, statut AS Statut, nom AS Nom, prenom AS 'Prénom', email AS Email,
                   telephone AS 'Téléphone', diplome AS 'Diplôme', specialite AS 'Spécialité',
                   departement AS 'Département', autres_informations AS 'Autres informations',
                   source_document AS 'Source document', feuille_section AS 'Feuille/Section',
                   date_import AS 'Date import', date_maj AS 'Date MAJ'
            FROM enseignants ORDER BY nom, prenom
        """, conn)


def load_teacher_history(teacher_id: int, db_path: Path = DB_PATH) -> pd.DataFrame:
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        return pd.read_sql_query("""
            SELECT field_name AS Champ, field_value AS Valeur, source_document AS Source,
                   feuille_section AS Section, observed_at AS 'Observé le'
            FROM field_history WHERE teacher_id=? ORDER BY id DESC
        """, conn, params=(int(teacher_id),))


def find_possible_duplicates(db_path: Path = DB_PATH) -> pd.DataFrame:
    base = load_teachers(db_path)
    if base.empty:
        return pd.DataFrame()
    rows = []
    records = base.to_dict("records")
    for i in range(len(records)):
        a = records[i]
        na = _norm(f"{a.get('Nom','')} {a.get('Prénom','')}")
        if not na:
            continue
        for j in range(i + 1, len(records)):
            b = records[j]
            nb = _norm(f"{b.get('Nom','')} {b.get('Prénom','')}")
            if not nb:
                continue
            score = SequenceMatcher(None, na, nb).ratio()
            if 0.86 <= score < 1.0:
                rows.append({
                    "ID 1": a["id"], "Enseignant 1": f"{a.get('Nom','')} {a.get('Prénom','')}",
                    "ID 2": b["id"], "Enseignant 2": f"{b.get('Nom','')} {b.get('Prénom','')}",
                    "Similarité": round(score, 3),
                })
    return pd.DataFrame(rows).sort_values("Similarité", ascending=False) if rows else pd.DataFrame()


def delete_teacher(teacher_id: int, db_path: Path = DB_PATH):
    ensure_db(db_path)
    with sqlite3.connect(db_path) as conn:
        conn.execute("DELETE FROM field_history WHERE teacher_id=?", (int(teacher_id),))
        conn.execute("DELETE FROM rh_conflicts WHERE teacher_id=?", (int(teacher_id),))
        conn.execute("DELETE FROM conflict_resolution_rules WHERE teacher_id=?", (int(teacher_id),))
        conn.execute("DELETE FROM enseignants WHERE id=?", (int(teacher_id),))
        conn.commit()


def export_teachers_excel(df: pd.DataFrame) -> bytes:
    output = io.BytesIO()
    conflicts = load_conflicts("all")
    mappings = load_unknown_headers()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        df.to_excel(writer, index=False, sheet_name="Enseignants RH")
        if not conflicts.empty:
            conflicts.to_excel(writer, index=False, sheet_name="Conflits RH")
        if not mappings.empty:
            mappings.to_excel(writer, index=False, sheet_name="Colonnes inconnues")
        for sheet_name, ws in writer.sheets.items():
            ws.freeze_panes(1, 0)
            sheet_df = df if sheet_name == "Enseignants RH" else (conflicts if sheet_name == "Conflits RH" else mappings)
            for idx, col in enumerate(sheet_df.columns):
                vals = sheet_df[col].fillna("").head(200) if col in sheet_df.columns else []
                width = min(max([len(str(col)) + 2] + [len(str(v)) + 2 for v in vals]), 50)
                ws.set_column(idx, idx, width)
    return output.getvalue()
